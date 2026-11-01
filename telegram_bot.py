import asyncio
import io
import logging
import os
import re
import secrets
import string
import sys
import threading
import time
import uuid
from datetime import datetime, timedelta, timezone

from dotenv import load_dotenv
from google import genai
from google.genai import types
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update, WebAppInfo
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

from check_stock import _norm_ar
from check_stock import check_stock as db_check_stock
from check_stock import escape_like, supabase
from orders import cancel_order, update_order_shipping, update_order_status
from orders import is_shipping_enabled, set_shipping_enabled
from subscriptions import (
    add_or_renew_subscription,
    check_and_update_subscriptions,
    get_tenant_subscription,
)
from merchant_settings import (
    get_merchant_channels,
    set_merchant_channels,
    normalize_channel_name,
    SUPPORTED_CHANNELS,
    CHANNEL_ALIASES,
)

sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
load_dotenv()

# حماية التوكنات في سجلات التطبيق
logging.basicConfig(level=logging.WARNING)
logging.getLogger("httpx").setLevel(logging.WARNING)

BOT_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
MODEL = "gemini-3.5-flash-lite"
BOT_USERNAME = os.environ.get("BOT_USERNAME", "ORIK_1_bot").lstrip("@")

if "SUPER_ADMIN_CHAT_ID" not in os.environ or not os.environ["SUPER_ADMIN_CHAT_ID"].strip():
    raise ValueError("⚠️ المتغير SUPER_ADMIN_CHAT_ID إجباري ويجب ضبطه في ملف .env لحماية النظام!")
SUPER_ADMIN_CHAT_ID = int(os.environ["SUPER_ADMIN_CHAT_ID"].strip())
SUPER_ADMIN_IDS = {SUPER_ADMIN_CHAT_ID}

SYSTEM_PROMPT = """أنت مساعد مخزن لصاحب محل ملابس مصري، بتتكلم معاه على تليجرام.
أسلوبك: لهجة مصرية بسيطة، ردود قصيرة جداً.

قواعد:
1. لأي استعلام عن المخزون استخدم check_stock. ممنوع تخمّن كميات.
2. لو صاحب المحل قال إنه باع، أو وصلته بضاعة، أو في مرتجع: حدد الصنف بـ check_stock عشان تجيب كود الصنف (sku)، وبعدين نادِ propose_stock_change.
3. propose_stock_change مابتنفذش حاجة، هي بس بتجهّز العملية. الأزرار هي اللي هتخلي صاحب المحل يأكد. ماتقولش إن العملية اتنفذت أبداً.
4. لو المقاس أو اللون ناقص أو في أكتر من احتمال، اسأل بدل ما تخمّن.
5. لو الكمية ماتقالتش افترض 1.
6. أنواع العمليات: sale = بيع من المحل (بينقص)، restock = وصلت بضاعة (بتزيد)، return = مرتجع (بيزيد).
7. لو صاحب المحل طلب إضافة منتج جديد (مش موجود قبل كده في المخزن) واكتملت بياناته (الاسم، السعر، المقاس، اللون، الكمية): استدعي دالة propose_new_product فوراً في الخلفية! ممنوع تقول باللسان 'تأكد من الأزرار' أو 'جهزت الإضافة' من غير ما تستدعي الدالة فعلياً.
8. لو ذكر كذا لون أو مقاس لنفس المنتج، نادِ propose_new_product لكل لون/مقاس في نفس الرد فوراً لتجهيزهم معاً دفعة واحدة!
9. تذكّر سياق الرسائل السابقة واجمع البيانات الناقصة من كلام التاجر بدون ما تعيد السؤال عن حاجات قالها بالفعل.
10. لو صاحب المحل طلب إضافة فرع أو عنوان للمحل (مثلاً: فرع مدينة نصر، شارع كذا، لوكيشن ماب كذا)، استخرج البيانات ونادِ propose_new_branch.
11. لو صاحب المحل طلب عمل كود خصم أو عرض ترويجي (مثلاً: كود EID20 خصم 20%، أو عرض شحن مجاني لأي حد يطلب قطعتين أو أكتر)، استخرج البيانات ونادِ propose_new_promo.
13. لو صاحب المحل بعت صورة منتج (أو كان بعت صورة من دقايق وبيكمل كلامه): الصورة بتتحفظ وبتتربط تلقائياً بالمنتج أول ما يضغط تأكيد. بص على الصورة والكلام اللي معاها وحدد نوع القطعة واللون. لو قال وصل/جاي/ضيف منتج جديد استخدم propose_new_product (الاسم، السعر، المقاس، اللون، الكمية). لو عايز يغيّر أو يستبدل صورة منتج موجود بالفعل (لأنه بعت صورة غلط أو مش عاجباه) استخدم propose_set_product_image فوراً باسم المنتج. لو السعر أو المقاس ناقص اسأل، ممنوع تخمّن. لو طلب تغيير صورة منتج ومكنش باعت صورة، اطلب منه يبعت الصورة الجديدة دلوقتي وقوله إنه يقدر يستخدم أمر /set_image لاختيار الصنف بسهولة. لو الصورة من غير أي كلام اسأله: المنتج ده إيه وسعره ومقاسه وكميته؟
12. اتفاقية حتمية: أي عملية بتغيّر أوردر أو مخزون لازم تعدي على أزرار التأكيد وماتتنفذش فوراً! لو صاحب المحل طلب إلغاء أوردر أو تعديل عنوان/موبايل/حالة أوردر، استدعِ propose_cancel_order أو propose_update_order لتجهيزها، والأزرار هي اللي هتأكد.
14. لو صاحب المحل سأل 'المخزن فيه إيه؟' أو 'عايز كشف المخزون' أو 'إيه اللي عندي في المخزن؟': نادِ list_inventory واديله كشف سريع بالأصناف والكميات المتاحة وقوله إنه يقدر يكتب /stock عشان يشوف الكشف كامل ومفصل.

معلومات: المقاسات M و L و XL و 40 و 41 و 42 و 43 وغيرها (لارج = L، ميديم = M، إكس لارج = XL). الألوان بالعربي: أسود، أبيض، رمادي، كحلي، بني، بيج، أحمر، أصفر، بنفسجي، أخضر.
"""

NOT_OWNER_MESSAGE = "البوت ده خاص بصاحب المحل. لو إنت صاحب المحل، افتح رابط الربط اللي وصلك."


def generate_sku() -> str:
    """توليد كود صنف بحروف لاتينية فقط (P + 6 حروف عشوائية)"""
    letters = "".join(secrets.choice(string.ascii_uppercase) for _ in range(6))
    return f"P{letters}"


def get_owner_tenant(chat_id):
    """بيرجّع التاجر لو الشات ده مربوط بأي تاجر في السيستم، غير كده None"""
    rows = (
        supabase.table("tenants")
        .select("id,name")
        .eq("telegram_chat_id", chat_id)
        .execute()
        .data
    )
    if rows:
        return rows[0]
    
    # دعم السوبر أدمن أو متجر ORIK الافتراضي
    super_admin_id = os.environ.get("SUPER_ADMIN_CHAT_ID", "").strip()
    if super_admin_id and str(chat_id).strip() == super_admin_id:
        def_rows = supabase.table("tenants").select("id,name").eq("name", "ORIK").limit(1).execute().data
        if not def_rows:
            def_rows = supabase.table("tenants").select("id,name").order("created_at", desc=False).limit(1).execute().data
        if def_rows:
            return def_rows[0]
            
    return None


# تعريف الاسم البديل لتفادي خطأ NameError
get_merchant_by_telegram = get_owner_tenant


# سجل محاولات الربط مع سقف أقصى لمنع استهلاك الذاكرة
LINK_CODE_ATTEMPTS = {}
MAX_LINK_CODE_ATTEMPTS_ENTRIES = 1000
MAX_FAILED_LINK_ATTEMPTS = 5
LINK_LOCKOUT_SECONDS = 15 * 60  # 15 دقيقة حظر بعد 5 محاولات فاشلة
LINK_CODE_RE = re.compile(r"^[A-F0-9]{16}_[0-9]{9,11}$")


def get_or_create_link_attempt(chat_id: int):
    now = time.time()
    if len(LINK_CODE_ATTEMPTS) > MAX_LINK_CODE_ATTEMPTS_ENTRIES:
        expired = [k for k, v in LINK_CODE_ATTEMPTS.items() if now > v.get("blocked_until", 0) and v.get("failures", 0) == 0]
        for k in expired:
            LINK_CODE_ATTEMPTS.pop(k, None)
        if len(LINK_CODE_ATTEMPTS) > MAX_LINK_CODE_ATTEMPTS_ENTRIES:
            for k in list(LINK_CODE_ATTEMPTS.keys())[:MAX_LINK_CODE_ATTEMPTS_ENTRIES // 4]:
                LINK_CODE_ATTEMPTS.pop(k, None)
    return LINK_CODE_ATTEMPTS.setdefault(chat_id, {"failures": 0, "blocked_until": 0})


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.effective_message
    chat_id = update.effective_chat.id
    code = context.args[0].strip().upper() if context.args else None

    if code:
        now = time.time()
        rate_info = get_or_create_link_attempt(chat_id)

        # 1. فحص الحظر المؤقت
        if rate_info["blocked_until"] > now:
            remaining_mins = int((rate_info["blocked_until"] - now) // 60) + 1
            await msg.reply_text(f"⛔ تم حظر المحاولات مؤقتاً لتكرار إدخال أكواد خاطئة. جرّب بعد {remaining_mins} دقيقة.")
            return

        def handle_failure(reason_text="الكود غلط أو اتستخدم قبل كده."):
            rate_info["failures"] += 1
            if rate_info["failures"] >= MAX_FAILED_LINK_ATTEMPTS:
                rate_info["blocked_until"] = now + LINK_LOCKOUT_SECONDS
                rate_info["failures"] = 0
                return "⛔ تم إدخال كود غير صالح 5 مرات. تم حظر محاولاتك لمدة 15 دقيقة لحماية الحسابات."
            else:
                remaining = MAX_FAILED_LINK_ATTEMPTS - rate_info["failures"]
                return f"{reason_text}\nباقي لك {remaining} محاولات قبل الحظر المؤقت."

        # فحص تطابق نمط الكود قبل البحث
        if not LINK_CODE_RE.fullmatch(code):
            await msg.reply_text(handle_failure())
            return

        # 2. البحث عن الكود
        rows = (
            supabase.table("tenants")
            .select("id,name,telegram_link_code,created_at")
            .eq("telegram_link_code", code)
            .execute()
            .data
        )
        if not rows:
            await msg.reply_text(handle_failure())
            return

        target_tenant = rows[0]
        stored_code = target_tenant.get("telegram_link_code") or ""

        # 3. فحص صلاحية الكود (24 ساعة)
        is_expired = False
        if "_" in stored_code:
            _, _, expiry_str = stored_code.rpartition("_")
            try:
                expiry_ts = float(expiry_str)
                if now > expiry_ts:
                    is_expired = True
            except ValueError:
                pass
        elif target_tenant.get("created_at"):
            try:
                created_dt = datetime.fromisoformat(target_tenant["created_at"].replace("Z", "+00:00"))
                if datetime.now(timezone.utc) - created_dt > timedelta(hours=24):
                    is_expired = True
            except Exception:
                pass

        if is_expired:
            supabase.table("tenants").update({"telegram_link_code": None}).eq("id", target_tenant["id"]).execute()
            await msg.reply_text(handle_failure("⛔ كود الربط انتهت صلاحيته (الصلاحية 24 ساعة فقط). اطلب من الأدمن توليد كود جديد."))
            return

        # الربط ناجح
        LINK_CODE_ATTEMPTS.pop(chat_id, None)
        supabase.table("tenants").update(
            {"telegram_chat_id": chat_id, "telegram_link_code": None}
        ).eq("id", target_tenant["id"]).execute()

        await msg.reply_text(
            f"تم ربط البوت بـ «{target_tenant['name']}» ✅\n"
            "هيوصلك هنا إشعار بكل أوردر جديد، وتقدر تدير مخزنك ومبيعاتك من هنا مباشرة.\n"
            "اكتب /help عشان تشوف تقدر تعمل إيه."
        )
        return

    tenant = get_owner_tenant(chat_id)
    if tenant:
        await msg.reply_text(f"أهلاً بيك في بوت «{tenant['name']}» 👋\nاكتب /help.")
    else:
        await msg.reply_text(NOT_OWNER_MESSAGE)


async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    tenant = get_owner_tenant(chat_id)
    is_admin = (chat_id == SUPER_ADMIN_CHAT_ID)

    if not tenant and not is_admin:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    text = ""
    if is_admin:
        text += (
            "👑 *لوحة تحكم مدير النظام (Super Admin):*\n"
            "• `/new_tenant <اسم المحل> [الموبايل] [1m|3m] [noship]` — إنشاء تاجر جديد (noship = بدون شحن)\n"
            "• `/shipping on|off <الرقم أو الاسم>` — تشغيل/إيقاف الشحن لتاجر\n"
            "• `/tenants` — عرض قائمة التجار وحالة اشتراكاتهم\n"
            "• `/renew <الرقم أو الاسم> [1m|3m|6m|1y]` — تجديد اشتراك تاجر وتفعيل البوت\n"
            "• `/channels <الرقم> [fb,ig,wa,tt]` — إدارة قنوات السوشيال ميديا للتاجر وكوتة الواتساب\n"
            "• `/pause_bot <الرقم>` — إيقاف بوت تاجر مؤقتاً\n"
            "• `/resume_bot <الرقم>` — إعادة تشغيل بوت تاجر\n"
            "• `/delete_tenant <id8> [confirm]` — حذف تاجر بأول 8 حروف من معرّفه نهائياً\n"
            "• `/link_page <الرقم> fb=<id> token=<token>` — ربط قنوات السوشيال ميديا\n\n"
            "━━━━━━━━━━━━━━━━━━━━\n"
        )

    text += (
        "📦 *إدارة المخزن والمبيعات:*\n"
        "تقدر تكلمني بالعامية المصرية عادي جداً (كتابة أو فويس)، مثلاً:\n"
        "• *بعت 2 تيشيرت أسود L*\n"
        "• *وصل 20 هودي رمادي L*\n"
        "• *المتاح إيه من التيشيرت الأبيض؟*\n"
        "• *الغي أوردر 5*\n"
        "• *عدل عنوان أوردر 5 لمصر الجديدة وموبايل 01012345678*\n"
        "• *ضيف فرع جديد في مدينة نصر...*\n"
        "• *اعمل كود خصم EID20 خصم 20%*\n\n"
        "📊 *الأوامر السريعة:*\n"
        "• `/dashboard` — فتح لوحة تحكم التاجر وإحصائيات المخزن والمبيعات (Web App)\n"
        "• `/stock` — عرض كشف المخزون المتاح والكميات الحالية لكل صنف\n"
        "• `/orders` — آخر 5 أوردرات وتفاصيلها\n"
        "• `/today` — تقرير ملخص مبيعات اليوم\n"
        "• `/cancel <رقم_الأوردر>` — إلغاء أوردر وإرجاع بضاعته للمخزن\n"
        "• `/shipping` — حالة الشحن، و`/shipping off` لو محلك استلام من المحل بس، و`/shipping on` لتشغيله\n"
        "• `/channels` — قنوات السوشيال ميديا المفعلة في باقتك (فيس، إنستا، تيك توك، واتساب)\n"
        "• `/set_image [اسم_الصنف]` — تغيير أو إضافة أو استبدال صورة لمنتج في المخزن\n\n"
        "📁 *إضافة وتحديث المخزن:*\n"
        "• *بصورة المنتج:* ابعت صورة الموديل واكتب معاه كابشن: «وصل 50 ترنج شتوي بني مقاس L بـ600» وهسجله بالصورة فوراً!\n"
        "• *بشيت إكسيل:* ابعت أي ملف Excel (.xlsx) أو CSV فيه بضاعتك، وهقرأه وأحدث المخزن فوراً!"
    )
    await update.message.reply_text(text, parse_mode="Markdown")


async def dashboard_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    tenant = get_owner_tenant(chat_id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    base_url = os.environ.get("DASHBOARD_URL") or os.environ.get("BASE_URL") or "http://localhost:8000"
    base_url = base_url.rstrip("/")
    dashboard_url = f"{base_url}/dashboard?tenant_id={tenant['id']}"

    if dashboard_url.startswith("https://"):
        keyboard = InlineKeyboardMarkup([[
            InlineKeyboardButton("📊 فتح لوحة التحكم (Web App)", web_app=WebAppInfo(url=dashboard_url))
        ]])
    else:
        keyboard = InlineKeyboardMarkup([[
            InlineKeyboardButton("📊 فتح لوحة التحكم", url=dashboard_url)
        ]])

    await update.message.reply_text(
        f"📊 *لوحة تحكم التاجر — «{tenant['name']}»*\n\n"
        "تقدر تشوف إحصائيات المحل، وتعدّل رابط شيت المخزن، وتدير صفحة الفيسبوك، وتوقف أو تشغل البوت من اللوحة:\n"
        f"{dashboard_url}",
        reply_markup=keyboard,
        parse_mode="Markdown"
    )


async def cancel_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tenant = get_owner_tenant(update.effective_chat.id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return
    if not context.args:
        await update.message.reply_text("اكتب رقم الأوردر بعد الأمر، مثلاً:\n/cancel 1")
        return

    order_num = context.args[0].strip().replace("#", "")
    res = cancel_order(tenant["id"], order_num)
    if res["status"] == "ok":
        await update.message.reply_text(f"تم إلغاء الأوردر #{order_num} ورجعت الكميات للمخزن بنجاح ✅")
    else:
        await update.message.reply_text(f"فشل الإلغاء: {res.get('reason', 'حصل خطأ')}")


async def new_tenant_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    if chat_id != SUPER_ADMIN_CHAT_ID:
        await update.message.reply_text("⛔ عذراً، الأمر ده مخصص للـ Super Admin فقط.")
        return

    if not context.args:
        await update.message.reply_text(
            "اكتب اسم المحل بعد الأمر، مثلاً:\n"
            "`/new_tenant محل السعادة`\n"
            "أو مع رقم الموبايل:\n"
            "`/new_tenant محل السعادة 01012345678`\n"
            "ولو المحل مفيهوش شحن (استلام من المحل بس) زوّد `noship`:\n"
            "`/new_tenant محل السعادة 01012345678 noship`",
            parse_mode="Markdown"
        )
        return

    full_text = " ".join(context.args)

    # علامة "بدون شحن": المحل بيبيع استلام من المحل بس
    noship = bool(re.search(r"(?<!\S)(noship|بدون_شحن|بدون-شحن)(?!\S)", full_text, re.IGNORECASE))
    if noship:
        full_text = re.sub(r"(?<!\S)(noship|بدون_شحن|بدون-شحن)(?!\S)", "", full_text, flags=re.IGNORECASE).strip()

    months = 1
    dur_match = re.search(r"\b(\d+)\s*(m|شهر|أشهر|شهور)\b|\b(1y|سنة)\b", full_text, re.IGNORECASE)
    if dur_match:
        matched_str = dur_match.group(0).lower()
        if "y" in matched_str or "سنة" in matched_str:
            months = 12
        elif dur_match.group(1):
            months = max(1, int(dur_match.group(1)))
        full_text = re.sub(r"\b\d+\s*(m|شهر|أشهر|شهور)\b|\b(1y|سنة)\b", "", full_text, flags=re.IGNORECASE).strip()

    phone_match = re.search(r"(\+?20\d{10}|01[0125]\d{8})", full_text.replace(" ", ""))
    if phone_match:
        phone = phone_match.group(0)
        name = re.sub(r"(\+?20[\s\d]{10,}|01[0125][\s\d]{8,})", "", full_text).strip()
    else:
        phone = None
        name = full_text.strip()

    msg = update.effective_message
    if not name:
        await msg.reply_text("لازم تكتب اسم المحل!")
        return

    raw_code = secrets.token_hex(8).upper()
    expiry_ts = int(time.time() + 24 * 3600)
    link_code = f"{raw_code}_{expiry_ts}"

    try:
        existing = (
            supabase.table("tenants")
            .select("id, name, telegram_chat_id, telegram_link_code")
            .ilike("name", escape_like(name))
            .execute()
            .data
        )
        if existing:
            ex = existing[0]
            if ex.get("telegram_chat_id"):
                await msg.reply_text(
                    f"⚠️ تنبيه: متجر «{name}» مسجل ومربوط بالفعل في السيستم!\n"
                    f"شوف بياناته عبر /tenants أو جدد اشتراكه بـ /renew {name} {months}m"
                )
                return
            else:
                code = link_code
                supabase.table("tenants").update({"telegram_link_code": link_code}).eq("id", ex["id"]).execute()
                deep_link = f"https://t.me/{BOT_USERNAME}?start={code}"
                await msg.reply_text(
                    f"ℹ️ المحل «{name}» متسجل مسبقاً وبانتظار التفعيل:\n\n"
                    f"🔑 كود الربط: {code}\n\n"
                    f"انسخ الرسالة دي وابعتها للتاجر:\n"
                    f"👇👇👇\n"
                    f"أهلاً بحضرتك في سيستم إدارة المبيعات والمخزن 🚀\n"
                    f"اضغط على الرابط التالي لتفعيل حسابك فوراً:\n"
                    f"{deep_link}\n\n"
                    f"أو ادخل على البوت @{BOT_USERNAME} واكتب:\n"
                    f"/start {code}"
                )
                return
    except Exception as e:
        print("Error checking existing tenant:", e)

    data = {
        "name": name,
        "bot_enabled": True,
        "telegram_link_code": link_code,
    }
    if phone:
        data["owner_phone"] = phone
    if noship:
        data["shipping_enabled"] = False

    try:
        res = supabase.table("tenants").insert(data).execute()
        if not res.data:
            await msg.reply_text("❌ حصل خطأ أثناء حفظ التاجر في قاعدة البيانات.")
            return

        t = res.data[0]
        sub = add_or_renew_subscription(t["id"], months=months)
        ends_str = sub["ends_at"].strftime("%Y-%m-%d")

        deep_link = f"https://t.me/{BOT_USERNAME}?start={link_code}"
        await msg.reply_text(
            f"🎉 تم إنشاء حساب التاجر بنجاح!\n\n"
            f"🏢 المحل: {t['name']}\n"
            f"🔑 كود الربط: {link_code}\n"
            f"📅 مدة الاشتراك: {months} شهر (ساري حتى: {ends_str})\n"
            f"🚚 الشحن: {'❌ غير متاح (استلام من المحل فقط)' if noship else '✅ متاح'}\n\n"
            f"انسخ الرسالة دي وابعتها للتاجر:\n"
            f"👇👇👇\n"
            f"أهلاً بحضرتك في سيستم إدارة المبيعات والمخزن 🚀\n"
            f"اضغط على الرابط التالي لتفعيل حسابك فوراً:\n"
            f"{deep_link}\n\n"
            f"أو ادخل على البوت @{BOT_USERNAME} واكتب:\n"
            f"/start {link_code}"
        )
    except Exception as e:
        hint = "\n(لو استخدمت noship شغّل docs/shipping_enabled.sql في Supabase الأول)" if noship else ""
        await msg.reply_text(f"❌ حدث خطأ: {e}{hint}")


async def link_page_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    if chat_id != SUPER_ADMIN_CHAT_ID:
        await update.message.reply_text("⛔ الأمر ده للـ Super Admin بس.")
        return

    if not context.args or len(context.args) < 2:
        await update.message.reply_text(
            "📋 *طريقة استخدام /link_page:*\n\n"
            "`/link_page <tenant_id> fb=<page_id>`\n"
            "`/link_page <tenant_id> ig=<instagram_id>`\n"
            "`/link_page <tenant_id> wa=<whatsapp_phone_number_id>`\n"
            "`/link_page <tenant_id> tiktok=<tiktok_account_id>`\n"
            "`/link_page <tenant_id> token=<fb_page_token>`\n"
            "`/link_page <tenant_id> wa_token=<wa_access_token>`\n"
            "`/link_page <tenant_id> tt_token=<tiktok_access_token>`\n\n"
            "ممكن تجمعهم في أمر واحد:\n"
            "`/link_page 3 fb=112233 ig=998877 wa=556644 token=EAABxx...`",
            parse_mode="Markdown"
        )
        return

    tenant_id_str = context.args[0].strip()
    target_clean = tenant_id_str.lower().replace("-", "")
    if len(target_clean) >= 8 and re.fullmatch(r"[0-9a-f\-]+", tenant_id_str.lower()):
        tenant, err = resolve_tenant_strict(tenant_id_str)
        if not tenant:
            await update.message.reply_text(err, parse_mode="Markdown")
            return
    else:
        tenant = resolve_tenant(tenant_id_str)
        if not tenant:
            await update.message.reply_text(f"❌ لم أجد تاجر بهذا الرقم أو الاسم: `{tenant_id_str}`\nاستخدم /tenants تشوف الأرقام الصح.", parse_mode="Markdown")
            return

    updates = {}
    args_lower = [a.lower() for a in context.args[1:]]
    if "unlink" in args_lower or "unlink_fb" in args_lower or "إلغاء" in args_lower:
        updates["facebook_page_id"] = None
        updates["facebook_page_token"] = None

    for arg in context.args[1:]:
        if "=" not in arg:
            continue
        key, _, val = arg.partition("=")
        key = key.lower().strip()
        val = val.strip()
        if not val:
            continue
        if val.lower() in ("none", "null", "del", "0", "clear", "حذف", "إلغاء"):
            val = None

        if key == "fb":
            updates["facebook_page_id"] = val
        elif key == "ig":
            updates["instagram_account_id"] = val
        elif key == "wa":
            updates["whatsapp_phone_number_id"] = val
        elif key == "token":
            updates["facebook_page_token"] = val
        elif key == "wa_token":
            updates["whatsapp_access_token"] = val
        elif key == "tiktok":
            updates["tiktok_account_id"] = val
        elif key == "tt_token":
            updates["tiktok_access_token"] = val

    if not updates:
        await update.message.reply_text("⚠️ مش فاهم أي معاملات! استخدم fb= أو ig= أو wa= أو token=")
        return

    try:
        supabase.table("tenants").update(updates).eq("id", tenant["id"]).execute()
    except Exception as e:
        await update.message.reply_text(f"❌ فشل التحديث: {e}")
        return

    lines = [f"✅ *تم ربط حسابات «{tenant['name']}» بنجاح!*\n"]
    field_names = {
        "facebook_page_id": "📘 فيسبوك Page ID",
        "instagram_account_id": "📸 إنستجرام Account ID",
        "whatsapp_phone_number_id": "🟢 واتساب Phone Number ID",
        "tiktok_account_id": "🎵 تيك توك Account ID",
        "facebook_page_token": "🔑 Facebook Page Token",
        "whatsapp_access_token": "🔑 WhatsApp Access Token",
        "tiktok_access_token": "🔑 TikTok Access Token",
    }
    for field, label in field_names.items():
        if field in updates:
            val = updates[field]
            display = "•••••••• (تم الحفظ بأمان)" if "token" in field.lower() else val
            lines.append(f"{label}: `{display}`")

    lines.append(f"\n🔗 *دلوقتي أي رسالة جاية للـ Page/Account ده هتروح لـ AI بتاع «{tenant['name']}» تلقائياً!*")
    await update.message.reply_text("\n".join(lines), parse_mode="Markdown")

    try:
        await update.message.delete()
    except Exception:
        pass


SHIPPING_ON_WORDS = {"on", "yes", "enable", "تشغيل", "شغل", "شغّل", "نعم", "متاح"}
SHIPPING_OFF_WORDS = {"off", "no", "disable", "ايقاف", "إيقاف", "وقف", "اقفل", "قفل", "لا", "بدون"}


async def shipping_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/shipping [on|off|<price>|<gov=price>] [التاجر] - إدارة وأسعار الشحن"""
    chat_id = update.effective_chat.id
    is_admin = chat_id == SUPER_ADMIN_CHAT_ID
    owner = await asyncio.to_thread(get_owner_tenant, chat_id)
    if not owner and not is_admin:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    mode = None
    new_flat_price = None
    custom_gov_rate = None
    rest = []

    for a in context.args or []:
        low = a.strip().lower()
        if mode is None and low in SHIPPING_ON_WORDS:
            mode = True
        elif mode is None and low in SHIPPING_OFF_WORDS:
            mode = False
        elif "=" in a:
            gov_part, _, p_part = a.partition("=")
            try:
                custom_gov_rate = (gov_part.strip(), float(p_part.strip()))
            except ValueError:
                rest.append(a)
        elif new_flat_price is None and re.match(r"^\d+(\.\d+)?$", a.strip()):
            new_flat_price = float(a.strip())
        else:
            rest.append(a)

    target = None
    if rest:
        if not is_admin:
            await update.message.reply_text("⛔ تحديد تاجر تاني للـ Super Admin بس. اكتب /shipping لمحلك.")
            return
        target, err = await asyncio.to_thread(resolve_tenant_strict, " ".join(rest))
        if not target:
            await update.message.reply_text(err)
            return
    elif owner:
        target = owner
    else:
        await update.message.reply_text("اكتب التاجر بعد الأمر، مثلاً: /shipping off 2 أو /shipping 65 2")
        return

    tid, name = target["id"], target["name"]

    # 1. إذا كان الأمر لتحديث سعر الشحن لمحافظة معينة
    if custom_gov_rate:
        c_gov, c_price = custom_gov_rate
        try:
            existing = supabase.table("shipping_rates").select("id").eq("tenant_id", tid).eq("governorate", c_gov).execute().data
            if existing:
                supabase.table("shipping_rates").update({"price": c_price}).eq("id", existing[0]["id"]).execute()
            else:
                supabase.table("shipping_rates").insert({"tenant_id": tid, "governorate": c_gov, "price": c_price}).execute()
            await asyncio.to_thread(set_shipping_enabled, tid, True)
            await update.message.reply_text(f"✅ تم تحديث سعر شحن محافظة «{c_gov}» في «{name}» إلى {c_price:g} ج.م بنجاح!")
            return
        except Exception as e:
            await update.message.reply_text(f"❌ فشل تحديث سعر الشحن: {e}")
            return

    # 2. إذا كان الأمر لتطبيق سعر شحن موحد
    if new_flat_price is not None:
        try:
            supabase.table("shipping_rates").delete().eq("tenant_id", tid).eq("governorate", "جميع المحافظات").execute()
            supabase.table("shipping_rates").insert({"tenant_id": tid, "governorate": "جميع المحافظات", "price": new_flat_price}).execute()
            await asyncio.to_thread(set_shipping_enabled, tid, True)
            await update.message.reply_text(
                f"✅ تم تطبيق الشحن الموحد لـ «{name}» بنجاح! 🚚\n\n"
                f"💰 سعر الشحن: {new_flat_price:g} ج.م لجميع المحافظات.\n"
                "الذكاء الاصطناعي وبوتات التواصل هيعتمدوا السعر الجديد فوراً في الردود والأوردرات."
            )
            return
        except Exception as e:
            await update.message.reply_text(f"❌ فشل تحديث سعر الشحن: {e}")
            return

    # 3. إذا كان الأمر لتشغيل أو إيقاف الشحن
    if mode is not None:
        ok, err = await asyncio.to_thread(set_shipping_enabled, tid, mode)
        if not ok:
            await update.message.reply_text(f"❌ ماقدرتش أغيّر الإعداد: {err}")
            return
        if mode:
            await update.message.reply_text(f"✅ تم تشغيل الشحن في «{name}» بنجاح.")
        else:
            await update.message.reply_text(
                f"✅ تم إيقاف الشحن في «{name}».\n"
                "من دلوقتي البوت هيقول للعملاء إن الاستلام من المحل بس."
            )
        return

    # 4. عرض حالة وأسعار الشحن الحالية
    enabled = await asyncio.to_thread(is_shipping_enabled, tid, True)
    state = "✅ متاح ومفعّل" if enabled else "❌ موقوف (استلام من المحل فقط)"

    rates_res = supabase.table("shipping_rates").select("governorate,price").eq("tenant_id", tid).execute()
    rates_list = rates_res.data or []

    if rates_list:
        rates_text = "\n📦 *أسعار الشحن المسجلة:*\n" + "\n".join(
            [f"  • {r['governorate']}: {float(r['price']):g} ج.م" for r in rates_list]
        )
    else:
        rates_text = "\n📦 *أسعار الشحن الحالية:* الافتراضي 50 ج.م لجميع المحافظات\n"

    await update.message.reply_text(
        f"🚚 *إعدادات الشحن في «{name}»:* {state}\n"
        f"{rates_text}\n"
        "💡 *أوامر التحكم السريع:*\n"
        "  • `/shipping 65` ⬅️ لتحديد سعر شحن موحد لكل المحافظات\n"
        "  • `/shipping القاهرة=40` ⬅️ لتحديد سعر لمحافظة معينة\n"
        "  • `/shipping on` أو `/shipping off` ⬅️ لتفعيل/إيقاف الشحن",
        parse_mode="Markdown"
    )


async def channels_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """التحكم في قنوات السوشيال ميديا المسموحة للتاجر (فيسبوك، إنستجرام، تيك توك، واتساب) وكوتة الواتساب"""
    chat_id = update.effective_chat.id
    owner = await asyncio.to_thread(get_owner_tenant, chat_id)
    is_admin = chat_id in SUPER_ADMIN_IDS

    if not owner and not is_admin:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    args = context.args or []
    target = owner
    channel_args = []

    # لو أدمن وبيحدد تاجر
    if is_admin:
        if args:
            possible_target = args[0]
            # فحص إذا كان أول وسيط هو تاجر
            t, _ = await asyncio.to_thread(resolve_tenant_strict, possible_target)
            if not t:
                t = await asyncio.to_thread(resolve_tenant, possible_target)
            if t:
                target = t
                channel_args = args[1:]
            else:
                channel_args = args
        elif not target:
            await update.message.reply_text("📋 حدد التاجر أولاً:\n`/channels <رقم أو اسم التاجر>`\nمثال: `/channels 1 fb,ig`", parse_mode="Markdown")
            return
    else:
        channel_args = args

    if not target:
        await update.message.reply_text("❌ لم أتمكن من تحديد التاجر.")
        return

    tid, name = target["id"], target["name"]

    # 1. لو في تعديل (للأدمن فقط)
    if channel_args:
        if not is_admin:
            await update.message.reply_text("⛔ تعديل باقات وقنوات الاشتراك متاح فقط لإدارة المنصة (Super Admin).")
            return

        arg_str = " ".join(channel_args).lower()
        cur_channels = await asyncio.to_thread(get_merchant_channels, tid, None, True)
        new_channels = dict(cur_channels)
        modified = False

        # أ) all / الكل
        if any(w in arg_str for w in ("all", "الكل", "كل")):
            for ch in SUPPORTED_CHANNELS:
                new_channels[ch] = True
            modified = True

        # ب) تحديد كوتة الواتساب: limit 1000 أو سقف 1000
        limit_match = re.search(r"(?:limit|سقف|باقة)\s*[:=]?\s*(\d+)", arg_str)
        if limit_match:
            new_channels["whatsapp_monthly_limit"] = int(limit_match.group(1))
            modified = True

        # ج) تصفير عداد الواتساب: reset أو تصفير
        if "reset" in arg_str or "تصفير" in arg_str:
            new_channels["whatsapp_used_conversations"] = 0
            modified = True

        # د) تعيين قائمة محددة مثل "fb,ig" أو "فيس وانستا"
        has_add_sub = any(c in arg_str for c in ("+", "-", "on", "off", "limit", "reset", "all", "الكل", "كل"))
        if not has_add_sub and not limit_match and "reset" not in arg_str:
            raw_tokens = re.split(r"[,+\s&]+", arg_str)
            found_channels = set()
            for tok in raw_tokens:
                norm = normalize_channel_name(tok)
                if norm in SUPPORTED_CHANNELS:
                    found_channels.add(norm)

            if found_channels:
                for ch in SUPPORTED_CHANNELS:
                    new_channels[ch] = (ch in found_channels)
                modified = True
        else:
            # تعديل نسبي (+ أو - أو on/off)
            for tok in channel_args:
                clean_tok = tok.strip().lower()
                if clean_tok.startswith("+"):
                    norm = normalize_channel_name(clean_tok[1:])
                    if norm in SUPPORTED_CHANNELS:
                        new_channels[norm] = True
                        modified = True
                elif clean_tok.startswith("-"):
                    norm = normalize_channel_name(clean_tok[1:])
                    if norm in SUPPORTED_CHANNELS:
                        new_channels[norm] = False
                        modified = True

            for ch in SUPPORTED_CHANNELS:
                for alias, real_ch in CHANNEL_ALIASES.items():
                    if real_ch == ch:
                        if f"{alias} on" in arg_str or f"{alias} تفعيل" in arg_str or f"تفعيل {alias}" in arg_str:
                            new_channels[ch] = True
                            modified = True
                        elif f"{alias} off" in arg_str or f"{alias} تعطيل" in arg_str or f"إيقاف {alias}" in arg_str:
                            new_channels[ch] = False
                            modified = True

        if modified:
            await asyncio.to_thread(set_merchant_channels, tid, new_channels)
            await update.message.reply_text(f"✅ تم تحديث قنوات اشتراك «{name}» بنجاح! 🎉")

    # 2. عرض الحالة الحالية للقنوات
    cur_channels = await asyncio.to_thread(get_merchant_channels, tid, None, True)

    fb_status = "🟢 مفعّل" if cur_channels.get("facebook") else "⚪ غير مفعّل"
    ig_status = "🟢 مفعّل" if cur_channels.get("instagram") else "⚪ غير مفعّل"
    tt_status = "🟢 مفعّل" if cur_channels.get("tiktok") else "⚪ غير مفعّل"

    wa_on = cur_channels.get("whatsapp")
    wa_used = cur_channels.get("whatsapp_used_conversations", 0)
    wa_limit = cur_channels.get("whatsapp_monthly_limit", 500)

    if wa_on:
        limit_text = f"({wa_used} / {wa_limit} محادثة)" if wa_limit > 0 else f"({wa_used} محادثة - بلا حدود)"
        wa_status = f"🟢 مفعّل {limit_text}"
    else:
        wa_status = "⚪ غير مفعّل (باقة إضافية)"

    help_admin_text = ""
    if is_admin:
        help_admin_text = (
            "\n━━━━━━━━━━━━━━━━━━━━\n"
            "⚙️ *أوامر تعديل الباقات والقنوات (للأدمن):*\n"
            "• تحديد قنوات الباقة: `/channels <التاجر> fb,ig`\n"
            "• تفعيل كل القنوات: `/channels <التاجر> all`\n"
            "• تفعيل/تعطيل قناة: `/channels <التاجر> +wa` أو `-tiktok`\n"
            "• تحديد باقة الواتساب: `/channels <التاجر> limit 1000`\n"
            "• تصفير عداد الواتساب: `/channels <التاجر> reset`"
        )

    await update.message.reply_text(
        f"📱 *قنوات السوشيال ميديا لمتجر «{name}»:*\n\n"
        f"• 📘 *فيسبوك (ماسينجر وكومنتات):* {fb_status}\n"
        f"• 📸 *إنستجرام (DM وكومنتات):* {ig_status}\n"
        f"• 🎵 *تيك توك (DM وكومنتات):* {tt_status}\n"
        f"• 🟢 *واتساب (WhatsApp API):* {wa_status}\n"
        f"{help_admin_text}",
        parse_mode="Markdown"
    )


async def autopost_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """التحكم في النشر التلقائي الذكي على فيسبوك: /autopost on/off أو /autopost 24"""
    chat_id = update.effective_chat.id
    owner = await asyncio.to_thread(get_owner_tenant, chat_id)
    is_admin = chat_id in SUPER_ADMIN_IDS

    if not owner and not is_admin:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    args = context.args or []
    target = owner
    mode = None
    interval = None

    for a in args:
        low = a.strip().lower()
        if low in ("on", "enable", "start", "شغال", "تفعيل"):
            mode = True
        elif low in ("off", "disable", "stop", "إيقاف", "تعطيل"):
            mode = False
        elif a.isdigit() and int(a) > 0:
            interval = int(a)
        elif is_admin and not target:
            target, _ = await asyncio.to_thread(resolve_tenant_strict, a)

    if not target:
        await update.message.reply_text("اكتب الأمر بهذا الشكل: /autopost on أو /autopost 24")
        return

    tid, name = target["id"], target["name"]
    t_row = supabase.table("tenants").select("autopost_enabled, autopost_interval_hours, autopost_last_at, facebook_page_id").eq("id", tid).execute().data
    row = t_row[0] if t_row else {}

    updates = {}
    if mode is not None:
        updates["autopost_enabled"] = mode
    if interval is not None:
        updates["autopost_interval_hours"] = interval

    if updates:
        supabase.table("tenants").update(updates).eq("id", tid).execute()
        cur_enabled = mode if mode is not None else row.get("autopost_enabled", False)
        cur_interval = interval if interval is not None else row.get("autopost_interval_hours", 24)
        await update.message.reply_text(
            f"✅ تم تحديث إعدادات النشر التلقائي في «{name}»:\n"
            f"• الحالة: {'🟢 مفعّل' if cur_enabled else '⚪ متوقف'}\n"
            f"• المعدل: كل {cur_interval} ساعة\n\n"
            "💡 يمكنك نشر بوست الآن فوراً بكتابة: /post_now"
        )
        return

    cur_enabled = row.get("autopost_enabled", False)
    cur_interval = row.get("autopost_interval_hours") or 24
    last_at = row.get("autopost_last_at") or "لم يتم النشر بعد"
    has_fb = bool(row.get("facebook_page_id"))

    fb_status = "✅ متصل بصفحة فيسبوك" if has_fb else "⚠️ غير مربوط بصفحة فيسبوك"

    await update.message.reply_text(
        f"📢 *إعدادات النشر التلقائي الذكي في «{name}»:*\n\n"
        f"• *الحالة:* {'🟢 مفعّل' if cur_enabled else '⚪ موقوف'}\n"
        f"• *معدل التكرار:* كل {cur_interval} ساعة\n"
        f"• *آخر بوست تم نشره:* {last_at}\n"
        f"• *حالة الصفحة:* {fb_status}\n\n"
        "💡 *الأوامر السريعة:*\n"
        "  • `/autopost on` ⬅️ تفعيل النشر التلقائي\n"
        "  • `/autopost off` ⬅️ إيقاف النشر التلقائي\n"
        "  • `/autopost 12` ⬅️ نشر بوست كل 12 ساعة\n"
        "  • `/post_now` ⬅️ نشر بوست إعلاني فوراً لأحد منتجات المخزن 🚀",
        parse_mode="Markdown"
    )


async def post_now_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """نشر بوست ترويجي فوراً على صفحة الفيسبوك من المخزون"""
    chat_id = update.effective_chat.id
    owner = await asyncio.to_thread(get_owner_tenant, chat_id)
    is_admin = chat_id in SUPER_ADMIN_IDS

    if not owner and not is_admin:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    target = owner
    if is_admin and context.args:
        target, err = await asyncio.to_thread(resolve_tenant_strict, " ".join(context.args))
        if not target:
            await update.message.reply_text(err)
            return

    if not target:
        await update.message.reply_text("حدد التاجر لنشر البوست.")
        return

    msg = await update.message.reply_text("⏳ جاري فحص المخزون واختيار أفضل قطعة وصياغة كابشن إعلاني جذاب ونشره على فيسبوك...")

    try:
        from server import trigger_autopost_for_tenant
        res = await trigger_autopost_for_tenant(target["id"])
        if res.get("status") == "success":
            p_name = res.get("product_name", "المنتج")
            post_id = res.get("post_id", "")
            await msg.edit_text(
                f"🚀 *تم نشر البوست بنجاح على فيسبوك!*\n\n"
                f"• *المنتج:* {p_name}\n"
                f"• *معرّف البوست:* `{post_id}`\n\n"
                "الذكاء الاصطناعي كتب كابشن بيعي جذاب يبرز الخامة والمقاسات والأسعار ومستعد للرد على العملاء وعمل لايك لتعليقاتهم تلقائياً! 📸✨",
                parse_mode="Markdown"
            )
        else:
            await msg.edit_text(f"❌ تعذر نشر البوست: {res.get('message', 'خطأ غير معروف')}")
    except Exception as e:
        await msg.edit_text(f"❌ حدث خطأ أثناء النشر: {e}")


async def offers_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """عرض وإدارة العروض الترويجية وكروت الهدايا: /offers"""
    chat_id = update.effective_chat.id
    owner = await asyncio.to_thread(get_owner_tenant, chat_id)
    is_admin = chat_id in SUPER_ADMIN_IDS

    if not owner and not is_admin:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    target = owner
    if is_admin and context.args:
        target, err = await asyncio.to_thread(resolve_tenant_strict, " ".join(context.args))
        if not target:
            await update.message.reply_text(err)
            return

    tid, name = target["id"], target["name"]
    res = supabase.table("promotions").select("*").eq("tenant_id", tid).order("created_at", desc=True).execute()
    promos = res.data or []

    if not promos:
        await update.message.reply_text(
            f"🎁 *العروض والخصومات في «{name}»:*\n\n"
            "مفيش أي عروض نشطة حالياً.\n\n"
            "💡 تقدر تضيف عرض جديد من لوحة التحكم أو تكلم البوت مباشرة، مثلاً:\n"
            "«اعمل عرض اشتري بـ 1000ج واحصل على فاوتشر 200ج للزيارة الجاية»\n"
            "أو «اعمل خصم 20% بكود SUMMER20»\n"
            "أو «اعمل شحن مجاني لأي طلب قطعتين أو أكثر»",
            parse_mode="Markdown"
        )
        return

    lines = [f"🎁 *العروض والخصومات النشطة في «{name}»:*\n"]
    for i, p in enumerate(promos, 1):
        ptype = p.get("promo_type", "percentage")
        val = p.get("discount_value", 0)
        code = f" (كود: `{p['code']}`)" if p.get("code") else " (تطبيق تلقائي)"
        min_v = p.get("min_order_value")
        min_i = p.get("min_order_items", 1)
        cond = f"شرط: أوردر بقيمة {min_v:g} ج.م" if min_v and min_v > 0 else (f"شرط: {min_i} قطع" if min_i > 1 else "بدون حد أدنى")
        target_str = f"\n   • 🎯 الأصناف: `{p['target_items']}`" if p.get("target_items") else ""

        type_lbl = "كارت هدايا / فاوتشر" if ptype in ("voucher", "gift_card") else ("شحن مجاني" if ptype == "free_shipping" else "خصم")
        lines.append(f"{i}. *{p.get('description', 'عرض')}*\n   • النوع: {type_lbl}{code}\n   • {cond}{target_str}")

    lines.append("\n💡 للتحكم الكامل أو إضافة/حذف أي عرض، افتح /dashboard")
    await update.message.reply_text("\n".join(lines), parse_mode="Markdown")


def resolve_tenant(identifier: str):
    """بحث مرن للتاجر بالرقم الترتيبي أو الاسم أو المعرّف (أو بادئة UUID فريدة 8+ حروف هيكس)."""
    rows = (
        supabase.table("tenants")
        .select("id, name, bot_enabled, telegram_chat_id, facebook_page_id, instagram_account_id, whatsapp_phone_number_id, tiktok_account_id")
        .order("name")
        .execute()
        .data or []
    )
    if not rows:
        return None

    clean = identifier.strip()
    if clean.isdigit():
        idx = int(clean) - 1
        if 0 <= idx < len(rows):
            return rows[idx]

    # 1. مطابقة UUID كامل
    for r in rows:
        if str(r.get("id", "")).lower() == clean.lower():
            return r

    # 2. مطابقة بادئة UUID فريدة (8+ حروف هيكس)
    clean_no_hyphens = clean.lower().replace("-", "")
    if len(clean_no_hyphens) >= 8 and re.fullmatch(r"[0-9a-f\-]+", clean.lower()):
        prefix_matches = [
            r for r in rows
            if str(r.get("id", "")).lower().startswith(clean.lower())
            or str(r.get("id", "")).lower().replace("-", "").startswith(clean_no_hyphens)
        ]
        if len(prefix_matches) == 1:
            return prefix_matches[0]

    # 3. مطابقة الاسم الجزئي
    for r in rows:
        if clean.lower() in str(r.get("name", "")).lower():
            return r

    return None


def resolve_tenant_strict(identifier: str):
    """
    مطابقة صارمة للتاجر (خاصة بعملية الحذف الحساسة):
    - تقبل UUID كامل أو بادئة لا تقل عن 8 خانات هيكس/UUID فريدة فقط.
    - ترفض الأسماء أو الأرقام الترتيبية أو البادئات الأقل من 8 خانات.
    - لو أكثر من تاجر يطابق نفس البادئة، ترفض العملية وتطلب تحديد المعرف بدقة.
    ترجع: (tenant_dict, error_message)
    """
    clean = (identifier or "").strip().lower()
    if not clean:
        return None, "يرجى كتابة معرّف التاجر (UUID كامل أو أول 8 حروف فريدة منه على الأقل)."

    clean_no_hyphens = clean.replace("-", "")
    if len(clean_no_hyphens) < 8 or not re.fullmatch(r"[0-9a-f\-]+", clean):
        return None, "❌ المعرّف غير صالح! يجب إدخال UUID كامل أو أول 8 حروف هيكس على الأقل (بدون أسماء أو أرقام ترتيبية)."

    rows = (
        supabase.table("tenants")
        .select("id, name, bot_enabled, telegram_chat_id")
        .execute()
        .data or []
    )
    if not rows:
        return None, "لا يوجد تجار مسجلين في السيستم."

    exact_matches = [r for r in rows if str(r.get("id", "")).lower() == clean]
    if len(exact_matches) == 1:
        return exact_matches[0], None

    prefix_matches = [
        r for r in rows
        if str(r.get("id", "")).lower().startswith(clean)
        or str(r.get("id", "")).lower().replace("-", "").startswith(clean_no_hyphens)
    ]
    if len(prefix_matches) == 1:
        return prefix_matches[0], None
    elif len(prefix_matches) > 1:
        matched_names = ", ".join(f"«{r['name']}» ({r['id'][:8]})" for r in prefix_matches)
        return None, f"⚠️ يوجد أكثر من تاجر يطابق هذا المعرّف ({matched_names}). يرجى كتابة معرّف أطول أو كامل لمنع الحذف بالخطأ."
    else:
        return None, f"❌ لم يتم العثور على أي تاجر بالمعرّف: `{clean}`."


async def list_tenants_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    if chat_id != SUPER_ADMIN_CHAT_ID:
        await update.message.reply_text("⛔ الأمر ده للـ Super Admin بس.")
        return

    rows = (
        supabase.table("tenants")
        .select("id,name,bot_enabled,facebook_page_id,instagram_account_id,whatsapp_phone_number_id,tiktok_account_id,telegram_chat_id")
        .order("name")
        .execute()
        .data
    )
    if not rows:
        await update.message.reply_text("مفيش تجار مسجلين لحد دلوقتي.")
        return

    lines = ["📋 *قائمة التجار المسجلين في السيستم:*\n"]
    for i, t in enumerate(rows, 1):
        is_active = t.get("bot_enabled", True)
        status_badge = "🟢 شغال" if is_active else "🔴 موقوف"

        sub = get_tenant_subscription(t["id"])
        if sub["has_subscription"]:
            if sub["is_active"]:
                sub_str = f"🟡 باقي {sub['days_left']} يوم ({sub['ends_at_str']})" if sub['days_left'] <= 3 else f"⏳ باقي {sub['days_left']} يوم ({sub['ends_at_str']})"
            else:
                sub_str = f"🔴 منتهي من {abs(sub['days_left'])} يوم"
        else:
            sub_str = "📅 غير محدد (جدده بـ /renew)"

        ch_settings = get_merchant_channels(t["id"])
        channels = []
        if t.get("facebook_page_id"):
            channels.append("📘 فيسبوك" if ch_settings.get("facebook") else "📘 فيسبوك (معطل)")
        if t.get("instagram_account_id"):
            channels.append("📸 إنستا" if ch_settings.get("instagram") else "📸 إنستا (معطل)")
        if t.get("tiktok_account_id"):
            channels.append("🎵 تيك توك" if ch_settings.get("tiktok") else "🎵 تيك توك (معطل)")
        if t.get("whatsapp_phone_number_id"):
            wa_badge = f"🟢 واتساب ({ch_settings.get('whatsapp_used_conversations', 0)}/{ch_settings.get('whatsapp_monthly_limit', 500)})" if ch_settings.get("whatsapp") else "🟢 واتساب (معطل)"
            channels.append(wa_badge)
        if t.get("telegram_chat_id"):
            channels.append("✈️ تليجرام")

        ch_str = " | ".join(channels) if channels else "لا توجد قنوات مربوطة بعد"

        lines.append(
            f"*{i}.* *«{t['name']}»* — {status_badge}\n"
            f"   📅 الاشتراك: {sub_str}\n"
            f"   🔗 القنوات: {ch_str}\n"
            f"   🆔 كود: `{t['id'][:8]}`\n"
        )

    lines.append("━━━━━━━━━━━━━━━━━━━━")
    lines.append("⚡ *التحكم السريع في التجار:*")
    lines.append("• تجديد اشتراك: `/renew <الرقم أو الاسم> [1m|3m|6m|1y]`")
    lines.append("• إدارة القنوات: `/channels <الرقم أو الاسم> [fb,ig,wa,tt]`")
    lines.append("• إيقاف مؤقت: `/pause_bot <الرقم>`")
    lines.append("• إعادة تشغيل: `/resume_bot <الرقم>`")
    lines.append("• حذف تاجر نهائياً: `/delete_tenant <id8>`")

    await update.message.reply_text("\n".join(lines), parse_mode="Markdown")


async def pause_bot_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    if chat_id != SUPER_ADMIN_CHAT_ID:
        await update.message.reply_text("⛔ الأمر ده للـ Super Admin بس.")
        return

    if not context.args:
        await update.message.reply_text("⚠️ اكتب رقم التاجر من قائمة /tenants أو اسمه:\nمثال: `/pause_bot 1`", parse_mode="Markdown")
        return

    target = " ".join(context.args).strip()
    target_clean = target.lower().replace("-", "")
    if len(target_clean) >= 8 and re.fullmatch(r"[0-9a-f\-]+", target.lower()):
        tenant, err = resolve_tenant_strict(target)
        if not tenant:
            await update.message.reply_text(err, parse_mode="Markdown")
            return
    else:
        tenant = resolve_tenant(target)
        if not tenant:
            await update.message.reply_text(f"❌ لم أجد تاجر بهذا الرقم أو الاسم: `{target}`", parse_mode="Markdown")
            return

    try:
        supabase.table("tenants").update({"bot_enabled": False}).eq("id", tenant["id"]).execute()
        await update.message.reply_text(
            f"⏸️ *تم إيقاف البوت للتاجر «{tenant['name']}» مؤقتاً!* 🔴\n\n"
            f"لإعادة تشغيله: `/resume_bot {tenant['name']}`",
            parse_mode="Markdown"
        )
    except Exception as e:
        await update.message.reply_text(f"❌ فشل إيقاف البوت: {e}")


async def resume_bot_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    if chat_id != SUPER_ADMIN_CHAT_ID:
        await update.message.reply_text("⛔ الأمر ده للـ Super Admin بس.")
        return

    if not context.args:
        await update.message.reply_text("⚠️ اكتب رقم التاجر من قائمة /tenants أو اسمه:\nمثال: `/resume_bot 1`", parse_mode="Markdown")
        return

    target = " ".join(context.args).strip()
    target_clean = target.lower().replace("-", "")
    if len(target_clean) >= 8 and re.fullmatch(r"[0-9a-f\-]+", target.lower()):
        tenant, err = resolve_tenant_strict(target)
        if not tenant:
            await update.message.reply_text(err, parse_mode="Markdown")
            return
    else:
        tenant = resolve_tenant(target)
        if not tenant:
            await update.message.reply_text(f"❌ لم أجد تاجر بهذا الرقم أو الاسم: `{target}`", parse_mode="Markdown")
            return

    try:
        supabase.table("tenants").update({"bot_enabled": True}).eq("id", tenant["id"]).execute()
        await update.message.reply_text(
            f"▶️ *تم إعادة تشغيل البوت للتاجر «{tenant['name']}» بنجاح!* 🟢",
            parse_mode="Markdown"
        )
    except Exception as e:
        await update.message.reply_text(f"❌ فشل تشغيل البوت: {e}")


def purge_tenant_data(tid: str):
    """حذف جميع بيانات التاجر بالترتيب الصحيح لتفادي أخطاء قيود المفاتيح الأجنبية (Foreign Key Constraints)"""
    # 1. جلب معرفات الأوردرات لحذف عناصرها
    order_rows = supabase.table("orders").select("id").eq("tenant_id", tid).execute().data or []
    order_ids = [r["id"] for r in order_rows]
    if order_ids:
        supabase.table("order_items").delete().in_("order_id", order_ids).execute()

    # 2. جلب معرفات الأصناف تحسباً لأي عناصر أوردرات تشير إليها
    variant_rows = supabase.table("variants").select("id").eq("tenant_id", tid).execute().data or []
    variant_ids = [r["id"] for r in variant_rows]
    if variant_ids:
        supabase.table("order_items").delete().in_("variant_id", variant_ids).execute()

    # 3. حذف حركات المخزون والأوردرات
    supabase.table("stock_movements").delete().eq("tenant_id", tid).execute()
    supabase.table("orders").delete().eq("tenant_id", tid).execute()

    # 4. حذف الرسائل والمحادثات
    conv_rows = supabase.table("conversations").select("id").eq("tenant_id", tid).execute().data or []
    conv_ids = [r["id"] for r in conv_rows]
    if conv_ids:
        supabase.table("messages").delete().in_("conversation_id", conv_ids).execute()
    supabase.table("conversations").delete().eq("tenant_id", tid).execute()

    # 5. حذف الجداول المساعدة (تقييمات، فروع، عروض، شحن، اشتراكات)
    for tbl in ["reviews", "branches", "promotions", "shipping_rates", "subscriptions"]:
        try:
            supabase.table(tbl).delete().eq("tenant_id", tid).execute()
        except Exception:
            pass

    # 6. حذف الأصناف ثم المنتجات ثم العملاء
    supabase.table("variants").delete().eq("tenant_id", tid).execute()
    supabase.table("products").delete().eq("tenant_id", tid).execute()
    supabase.table("customers").delete().eq("tenant_id", tid).execute()

    # 7. حذف التاجر نفسه نهائياً
    supabase.table("tenants").delete().eq("id", tid).execute()


async def delete_tenant_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """حذف التاجر وبياناته نهائياً بالتحقق الصارم من أول 8 حروف للـ UUID"""
    chat_id = update.effective_chat.id
    if chat_id != SUPER_ADMIN_CHAT_ID:
        await update.message.reply_text("⛔ الأمر ده للـ Super Admin بس.")
        return

    if not context.args:
        await update.message.reply_text(
            "⚠️ اكتب أول 8 حروف من معرّف التاجر (UUID) من قائمة /tenants:\n"
            "مثال: `/delete_tenant a1b2c3d4`",
            parse_mode="Markdown"
        )
        return

    args = context.args
    is_confirmed = False
    if len(args) > 1 and args[-1].lower() in ["confirm", "تأكيد", "نعم", "yes"]:
        is_confirmed = True
        target = " ".join(args[:-1]).strip()
    else:
        target = " ".join(args).strip()

    tenant, err = resolve_tenant_strict(target)
    if not tenant:
        await update.message.reply_text(err, parse_mode="Markdown")
        return

    tid = tenant["id"]
    tid8 = tid[:8]

    if not is_confirmed:
        orders_count = 0
        try:
            ord_res = supabase.table("orders").select("id", count="exact").eq("tenant_id", tid).execute()
            orders_count = ord_res.count if ord_res.count is not None else len(ord_res.data or [])
        except Exception:
            pass

        await update.message.reply_text(
            f"⚠️ تحذير: هل أنت متأكد من حذف التاجر «{tenant['name']}» نهائياً؟\n(المعرّف: `{tid}`)\n"
            f"ده هيحذف كل منتجاته ومخزنه وسجل أوردراته ({orders_count} أوردر) بالكامل ولا يمكن التراجع!\n\n"
            f"للتأكيد، اكتب:\n"
            f"`/delete_tenant {tid8} confirm`",
            parse_mode="Markdown"
        )
        return

    try:
        await asyncio.to_thread(purge_tenant_data, tid)
        await update.message.reply_text(
            f"🗑️ تم حذف التاجر «{tenant['name']}» وبياناته من السيستم نهائياً ✅\n(المعرّف: `{tid}`)",
            parse_mode="Markdown"
        )
    except Exception as e:
        await update.message.reply_text(f"❌ فشل الحذف: {e}")


async def renew_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    if chat_id != SUPER_ADMIN_CHAT_ID:
        await update.message.reply_text("⛔ الأمر ده للـ Super Admin بس.")
        return

    if not context.args:
        await update.message.reply_text("📋 طريقة التجديد:\n`/renew <رقم أو اسم أو كود التاجر> [المدة]`\nمثال: `/renew 1 3m`", parse_mode="Markdown")
        return

    target = context.args[0].strip()
    months = 1
    if len(context.args) > 1:
        m_str = context.args[1].lower().strip()
        if "3" in m_str:
            months = 3
        elif "6" in m_str:
            months = 6
        elif "12" in m_str or "y" in m_str or "سنة" in m_str:
            months = 12
        elif "شهرين" in m_str:
            months = 2

    target_clean = target.lower().replace("-", "")
    if len(target_clean) >= 8 and re.fullmatch(r"[0-9a-f\-]+", target.lower()):
        tenant, err = resolve_tenant_strict(target)
        if not tenant:
            await update.message.reply_text(err, parse_mode="Markdown")
            return
    else:
        tenant = resolve_tenant(target)
        if not tenant:
            await update.message.reply_text(f"❌ لم أجد تاجر بهذا الرقم أو الاسم: `{target}`\nاستخدم /tenants", parse_mode="Markdown")
            return

    tid = tenant["id"]
    res = add_or_renew_subscription(tid, months=months)
    supabase.table("tenants").update({"bot_enabled": True}).eq("id", tid).execute()

    ends_str = res["ends_at"].strftime("%Y-%m-%d")
    await update.message.reply_text(
        f"✅ *تم تجديد اشتراك «{tenant['name']}» بنجاح!* 🎉\n\n"
        f"📅 *مدة التجديد:* {months} شهر\n"
        f"⏳ *تاريخ الانتهاء الجديد:* `{ends_str}` (باقي {res['days_left']} يوم)\n"
        f"🟢 *حالة البوت:* شغال ويرد تلقائياً على كل الرسائل.",
        parse_mode="Markdown"
    )

    tg_chat = tenant.get("telegram_chat_id")
    if tg_chat:
        try:
            await context.bot.send_message(
                chat_id=tg_chat,
                text=(
                    f"🎉 مبروك! تم تجديد اشتراك متجرك «{tenant['name']}» بنجاح!\n\n"
                    f"📅 مدة التجديد: {months} شهر\n"
                    f"⏳ ساري حتى: {ends_str}\n\n"
                    "خدمة الرد التلقائي وإدارة الأوردرات والمخزن تعمل بكامل طاقتها 🚀"
                )
            )
        except Exception as e:
            print(f"Failed to notify tenant {tenant['name']} of renewal:", e)


async def check_subscriptions_and_notify(app):
    """فحص دوري لاشتراكات التجار وإرسال التنبيهات بنص عادي خالٍ من الرموز لتفادي أخطاء parse_mode"""
    alerts = check_and_update_subscriptions()

    for w in alerts.get("warnings", []):
        tid = w["tenant_id"]
        rows = supabase.table("tenants").select("id, name, telegram_chat_id").eq("id", tid).execute().data
        if not rows:
            continue
        tenant = rows[0]
        days = w["days_left"]
        ends_str = w["ends_at"].strftime("%Y-%m-%d")

        try:
            await app.bot.send_message(
                chat_id=SUPER_ADMIN_CHAT_ID,
                text=(
                    f"🔔 تنبيه تجديد اشتراك: ⏳\n\n"
                    f"المحل: «{tenant['name']}»\n"
                    f"متبقي على انتهاء الاشتراك: {days} أيام (ينتهي في: {ends_str})\n"
                    f"تواصل مع التاجر لتجديد الاشتراك.\n\n"
                    f"للتجديد السريع:\n/renew {tenant['id'][:8]} 1m أو 3m"
                )
            )
        except Exception as e:
            print("Failed to alert admin:", e)

        if tenant.get("telegram_chat_id"):
            try:
                await app.bot.send_message(
                    chat_id=tenant["telegram_chat_id"],
                    text=(
                        f"⚠️ تنبيه هام بخصوص اشتراكك:\n\n"
                        f"عزيزي صاحب متجر «{tenant['name']}»، متبقي {days} أيام فقط على انتهاء فترة اشتراكك في السيستم (ينتهي في: {ends_str}).\n\n"
                        "يرجى التواصل مع إدارة السيستم لتجديد الاشتراك وتجنب توقف خدمة الرد التلقائي على العملاء 🙏"
                    )
                )
            except Exception as e:
                print("Failed to alert tenant:", e)

    for exp in alerts.get("expirations", []):
        tid = exp["tenant_id"]
        rows = supabase.table("tenants").select("id, name, telegram_chat_id").eq("id", tid).execute().data
        if not rows:
            continue
        tenant = rows[0]
        ends_str = exp["ends_at"].strftime("%Y-%m-%d")

        supabase.table("tenants").update({"bot_enabled": False}).eq("id", tid).execute()

        try:
            await app.bot.send_message(
                chat_id=SUPER_ADMIN_CHAT_ID,
                text=(
                    f"🔴 انتهاء اشتراك وإيقاف البوت: ⏸️\n\n"
                    f"المحل: «{tenant['name']}»\n"
                    f"انتهى الاشتراك بتاريخ {ends_str} وتم إيقاف الرد الآلي للتاجر تلقائياً.\n\n"
                    f"للتجديد بعد الدفع:\n/renew {tenant['id'][:8]} 1m"
                )
            )
        except Exception as e:
            print("Failed to alert admin on expiration:", e)

        if tenant.get("telegram_chat_id"):
            try:
                await app.bot.send_message(
                    chat_id=tenant["telegram_chat_id"],
                    text=(
                        f"🔴 تنبيه: انتهت فترة اشتراك متجرك «{tenant['name']}»!\n\n"
                        "تم إيقاف خدمة الرد التلقائي وإدارة المبيعات مؤقتاً.\n"
                        "يرجى التواصل مع الإدارة وسداد الاشتراك لإعادة تفعيل الخدمة فوراً 🚀"
                    )
                )
            except Exception as e:
                print("Failed to alert tenant on expiration:", e)


async def orders_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tenant = get_owner_tenant(update.effective_chat.id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return
    rows = (
        supabase.table("orders")
        .select("order_number,customer_name,governorate,total,status")
        .eq("tenant_id", tenant["id"])
        .order("created_at", desc=True)
        .limit(5)
        .execute()
        .data
    )
    if not rows:
        await update.message.reply_text("لسه مفيش أوردرات.")
        return
    lines = [
        f"#{r['order_number']} | {r['customer_name']} | {r['governorate']} | {float(r['total']):.0f} جنيه | {r['status']}"
        for r in rows
    ]
    await update.message.reply_text("آخر الأوردرات:\n" + "\n".join(lines))


async def today_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tenant = get_owner_tenant(update.effective_chat.id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return
    since = (datetime.now(timezone.utc) - timedelta(hours=24)).isoformat()
    rows = (
        supabase.table("orders")
        .select("total,status")
        .eq("tenant_id", tenant["id"])
        .gte("created_at", since)
        .execute()
        .data
    )
    valid = [r for r in rows if r["status"] != "cancelled"]
    revenue = sum(float(r["total"]) for r in valid)
    await update.message.reply_text(
        f"آخر 24 ساعة:\n🛒 أوردرات من البوت: {len(valid)}\n💰 إجمالي: {revenue:.0f} جنيه"
    )


async def stock_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """عرض كشف المخزون المتاح والكميات الحالية لكل صنف في المحل مع حالة الصور"""
    tenant = get_owner_tenant(update.effective_chat.id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    rows = (
        supabase.table("variants")
        .select("sku, size, color, stock, reserved, available, image_url, products!inner(name, price, image_url)")
        .eq("tenant_id", tenant["id"])
        .eq("active", True)
        .execute()
        .data or []
    )
    if not rows:
        await update.message.reply_text("📦 المخزن فاضي حالياً! تقدر تضيف بضاعة فويس أو كتابة أو تبعت شيت إكسيل.")
        return

    lines = [f"📦 كشف مخزون «{tenant['name']}» المتاح حالياً:\n"]
    total_qty = 0
    for r in rows:
        pname = r.get("products", {}).get("name", "منتج")
        price = float(r.get("products", {}).get("price") or 0)
        avail = int(r.get("available") if r.get("available") is not None else (r.get("stock", 0) - r.get("reserved", 0)))
        total_qty += max(0, avail)
        has_img = "🖼️" if (r.get("products", {}).get("image_url") or r.get("image_url")) else "📷"
        lines.append(f"• {pname} ({r.get('color')} - {r.get('size')}): {avail} قطعة | {price:.0f}ج | كود: {r.get('sku')} {has_img}")

    lines.append(f"\n📊 إجمالي الأصناف: {len(rows)} صنف | {total_qty} قطعة في المخزن")
    lines.append("📸 لتغيير أو إضافة صورة لأي صنف: استخدم أمر /set_image أو ابعت الصورة واكتب (غير صورة اسم_الصنف).")
    await update.message.reply_text("\n".join(lines))


# جلسات المحادثة وإدارة الذاكرة مع قفل أمان
MERCHANT_SESSIONS = {}
MERCHANT_SESSIONS_LOCK = threading.Lock()
MAX_MERCHANT_SESSIONS = 200
MERCHANT_SESSION_TTL = 4 * 3600  # 4 ساعات


def run_llm(chat_id: int, content, tools: list):
    """
    تشغيل نموذج الذكاء الاصطناعي مع الأدوات وسياق المحادثة المحفوظ.
    يتم إنشاء Chat جديد لكل رسالة بناءً على السجل السابق (history)
    دون تعديل الخصائص الداخلية الخاصة مثل _config.
    """
    now = time.time()
    config = types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT, tools=tools)

    with MERCHANT_SESSIONS_LOCK:
        if len(MERCHANT_SESSIONS) > MAX_MERCHANT_SESSIONS:
            expired = [cid for cid, sess in MERCHANT_SESSIONS.items() if now - sess.get("last_active", 0) > MERCHANT_SESSION_TTL]
            for cid in expired:
                MERCHANT_SESSIONS.pop(cid, None)
            if len(MERCHANT_SESSIONS) > MAX_MERCHANT_SESSIONS:
                sorted_s = sorted(MERCHANT_SESSIONS.keys(), key=lambda k: MERCHANT_SESSIONS[k].get("last_active", 0))
                for k in sorted_s[:MAX_MERCHANT_SESSIONS // 5]:
                    MERCHANT_SESSIONS.pop(k, None)

        session = MERCHANT_SESSIONS.get(chat_id)
        history = list(session["history"]) if session and "history" in session else None

    # إنشاء chat جديد بسياق موثق
    chat = client.chats.create(model=MODEL, config=config, history=history)

    for attempt in range(3):
        try:
            resp = chat.send_message(content)
            raw_history = chat.get_history(curated=True)
            trimmed_history = raw_history[-20:] if raw_history else []
            clean_history = []
            for msg in trimmed_history:
                parts = [p for p in getattr(msg, "parts", []) if not getattr(p, "inline_data", None)]
                if parts:
                    msg.parts = parts
                    clean_history.append(msg)

            with MERCHANT_SESSIONS_LOCK:
                MERCHANT_SESSIONS[chat_id] = {
                    "history": clean_history,
                    "last_active": time.time(),
                }
            return resp.text or "تمام، أقدر أساعدك في إيه تاني؟"
        except Exception as e:
            msg = str(e)
            if "429" in msg:
                print("وصلنا لحد الحصة المجانية:", msg[:300])
                return "وصلنا لحد الاستخدام المجاني للذكاء الاصطناعي دلوقتي، جرّب بعد شوية."
            if "503" in msg and attempt < 2:
                time.sleep(3 * (attempt + 1))
                continue
            print("خطأ من الموديل:", type(e).__name__, msg[:300])
            with MERCHANT_SESSIONS_LOCK:
                MERCHANT_SESSIONS.pop(chat_id, None)
            return "حصلت مشكلة مؤقتة، جرّب تاني بعد شوية."


MAX_PENDING_OPERATIONS = 500
PENDING_OPERATION_TTL = 7200  # ساعتان


def _get_bot_data(context):
    if hasattr(context, "application") and hasattr(context.application, "bot_data"):
        return context.application.bot_data
    if hasattr(context, "bot_data"):
        return context.bot_data
    if isinstance(context, dict):
        return context
    return {}


def store_pending_operation(context, pid: str, data: dict):
    now = time.time()
    bot_data = _get_bot_data(context)
    pending = bot_data.setdefault("pending", {})
    if len(pending) > MAX_PENDING_OPERATIONS:
        expired = [k for k, v in pending.items() if now - v.get("created_at", 0) > PENDING_OPERATION_TTL]
        for k in expired:
            pending.pop(k, None)
        if len(pending) > MAX_PENDING_OPERATIONS:
            sorted_keys = sorted(pending.keys(), key=lambda k: pending[k].get("created_at", 0))
            for k in sorted_keys[:MAX_PENDING_OPERATIONS // 5]:
                pending.pop(k, None)
    data["created_at"] = now
    pending[pid] = data


def pop_pending_operation(context, pid: str) -> dict:
    bot_data = _get_bot_data(context)
    pending = bot_data.get("pending", {}).pop(pid, None)
    if not pending:
        return None
    if time.time() - pending.get("created_at", 0) > PENDING_OPERATION_TTL:
        return None
    return pending


# ---------------------------------------------------------------
# 🖼️ صور المنتجات: استقبال من تليجرام + رفع على Supabase Storage
# ---------------------------------------------------------------
IMAGES_BUCKET = os.environ.get("PRODUCT_IMAGES_BUCKET", "product-images")
MAX_IMAGE_BYTES = 5 * 1024 * 1024  # 5 ميجا
ALLOWED_IMAGE_MIMES = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}
LAST_PHOTOS = {}  # chat_id -> {"bytes", "mime", "ts"}  (آخر صورة بعتها التاجر)
LAST_PHOTO_TTL = 15 * 60  # 15 دقيقة
MAX_LAST_PHOTOS = 200


def remember_photo(chat_id: int, data: bytes, mime: str):
    now = time.time()
    if len(LAST_PHOTOS) > MAX_LAST_PHOTOS:
        for k in [k for k, v in LAST_PHOTOS.items() if now - v["ts"] > LAST_PHOTO_TTL]:
            LAST_PHOTOS.pop(k, None)
    LAST_PHOTOS[chat_id] = {"bytes": data, "mime": mime, "ts": now}


def get_fresh_photo(chat_id: int):
    ph = LAST_PHOTOS.get(chat_id)
    if ph and time.time() - ph["ts"] <= LAST_PHOTO_TTL:
        return ph
    LAST_PHOTOS.pop(chat_id, None)
    return None


def upload_product_image(tenant_id: str, data: bytes, mime: str) -> str:
    """تحسين وضغط الصورة ورفعها على Supabase Storage (bucket عام) وبيرجّع الرابط العام."""
    from image_utils import optimize_product_image
    # ضغط وتصغير وتحسين الصورة وتعديل تدوير الكاميرا لتسريع التحميل للعميل
    data, mime, ext = optimize_product_image(data, mime=mime)
    path = f"{tenant_id}/{uuid.uuid4().hex}.{ext}"
    bucket = supabase.storage.from_(IMAGES_BUCKET)
    try:
        bucket.upload(path, data, {"content-type": mime})
    except Exception as first_err:
        # أول مرة؟ جرّب تنشئ الـ bucket كعام وكرر الرفع
        try:
            supabase.storage.create_bucket(IMAGES_BUCKET, options={"public": True})
            bucket.upload(path, data, {"content-type": mime})
        except Exception:
            raise RuntimeError(
                f"تعذر رفع الصورة. اتأكد إن في bucket عام اسمه {IMAGES_BUCKET} في Supabase Storage ({type(first_err).__name__})"
            )
    url = bucket.get_public_url(path)
    if isinstance(url, dict):  # اختلاف بين إصدارات المكتبة
        url = url.get("publicURL") or url.get("publicUrl") or ""
    return str(url).rstrip("?")


def _find_product_by_name(tenant_id: str, name: str):
    rows = (
        supabase.table("products").select("id,name,image_url")
        .eq("tenant_id", tenant_id).execute().data or []
    )
    target = _norm_ar(name)
    for r in rows:
        if _norm_ar(r.get("name", "")) == target:
            return r
    for r in rows:
        rn = _norm_ar(r.get("name", ""))
        if target and (target in rn or rn in target):
            return r
    tokens = [t for t in target.split() if len(t) > 2]
    for r in rows:
        rn = _norm_ar(r.get("name", ""))
        if any(tok in rn for tok in tokens):
            return r
    return None


async def reply_with_proposals(update, context, chat_id, tenant_id, proposals, reply, intro="هنفّذ العملية دي:", image=None):
    """لو في عمليات مقترحة: اعرض الملخص بأزرار التأكيد، وإلا ابعت رد الموديل."""
    if not proposals:
        await update.message.reply_text(reply)
        return
    pid = uuid.uuid4().hex[:12]
    data = {"chat_id": chat_id, "tenant_id": tenant_id, "items": proposals}
    if image and any(p.get("type") in ("new_product", "set_image") for p in proposals):
        data["image"] = {"bytes": image["bytes"], "mime": image["mime"]}
        LAST_PHOTOS.pop(chat_id, None)
    store_pending_operation(context, pid, data)
    summary = build_proposals_summary(proposals)
    if data.get("image"):
        summary += "\n🖼️ هتتربط الصورة اللي بعتها بالمنتج."
    keyboard = InlineKeyboardMarkup([[
        InlineKeyboardButton("✅ تأكيد", callback_data=f"ok:{pid}"),
        InlineKeyboardButton("❌ إلغاء", callback_data=f"no:{pid}"),
    ]])
    await update.message.reply_text(f"{intro}\n{summary}\n\nتأكيد؟", reply_markup=keyboard)


def build_tools(tenant_id: str, proposals: list, has_image: bool = False) -> list:
    """
    تجهيز قائمة أدوات Gemini التفاعلية لتاجر معين مع docstrings كاملة ومفصلة.
    تُستخدم في handle_text و handle_voice دون تكرار كود.
    أي عملية تعدل المخزون أو الأوردرات لا تنفذ مباشرة بل تُضاف إلى proposals بانتظار التأكيد.
    """
    def check_stock(product_name: str, size: str = "", color: str = "") -> dict:
        """بتكشف المخزون لمنتج بمقاس ولون معين، وبترجّع كود الصنف (sku) والكمية المتاحة.

        Args:
            product_name: اسم المنتج أو جزء منه (تيشيرت، هودي، بنطلون)
            size: المقاس M أو L أو XL (فاضي لو ماتحددش)
            color: اللون بالعربي (فاضي لو ماتحددش)
        """
        return db_check_stock(
            tenant_id, product_name, size=size or None, color=color or None, reveal_qty=True
        )

    def list_inventory() -> dict:
        """بتعرض ملخص لكل الأصناف والكميات والأسعار المتاحة حالياً في مخزن المحل. استخدمها لما التاجر يسأل: 'المخزن فيه إيه؟' أو 'إيه اللي عندي في المخزن؟' أو 'عايز كشف المخزون'."""
        rows = (
            supabase.table("variants")
            .select("sku, size, color, stock, reserved, available, products!inner(name, price)")
            .eq("tenant_id", tenant_id)
            .eq("active", True)
            .execute()
            .data or []
        )
        if not rows:
            return {"status": "empty", "message": "المخزن فاضي حالياً"}
        items = []
        for r in rows:
            pname = r.get("products", {}).get("name", "منتج")
            price = float(r.get("products", {}).get("price") or 0)
            avail = int(r.get("available") if r.get("available") is not None else (r.get("stock", 0) - r.get("reserved", 0)))
            items.append({
                "product": pname,
                "color": r.get("color"),
                "size": r.get("size"),
                "available": avail,
                "price": price,
                "sku": r.get("sku"),
            })
        return {"status": "ok", "total_items": len(items), "inventory": items}

    def propose_stock_change(sku: str, qty: int, kind: str) -> dict:
        """بتجهّز تعديل على المخزون (بيع، بضاعة، مرتجع) وتستنى تأكيد صاحب المحل بالأزرار. مابتنفذش حاجة مباشرة.

        Args:
            sku: كود الصنف اللي جاي من check_stock
            qty: الكمية (رقم موجب)
            kind: نوع العملية: sale (بيع من المحل) أو restock (وصلت بضاعة) أو return (مرتجع)
        """
        if qty <= 0:
            return {"status": "error", "reason": "الكمية لازم تكون أكبر من صفر"}
        if kind not in ("sale", "restock", "return"):
            return {"status": "error", "reason": "نوع العملية غلط"}
        rows = (
            supabase.table("variants")
            .select("id,sku,size,color,available,products!inner(name)")
            .eq("tenant_id", tenant_id)
            .eq("sku", sku)
            .execute()
            .data
        )
        if not rows:
            return {"status": "error", "reason": f"الكود {sku} مش موجود"}
        v = rows[0]
        if kind == "sale" and v["available"] < qty:
            return {"status": "error", "reason": f"المتاح {v['available']} بس"}
        label = f"{v['products']['name']} {v['color']} {v['size']}"
        change = -qty if kind == "sale" else qty
        proposals.append(
            {"type": "movement", "variant_id": v["id"], "label": label, "change": change, "kind": kind}
        )
        return {"status": "waiting_confirmation", "item": label, "change": change}

    def propose_new_product(name: str, price: float, size: str, color: str, qty: int = 1, material: str = "") -> dict:
        """بتجهّز إضافة منتج أو صنف جديد خالص في المخزن وتستنى تأكيد صاحب المحل بالأزرار.

        Args:
            name: اسم المنتج (مثلاً: جاكيت جينز، سويت شيرت)
            price: سعر البيع بالجنيه
            size: المقاس (M, L, XL, 42, إلخ)
            color: اللون بالعربي (أسود، كحلي، إلخ)
            qty: الكمية المتاحة الأولية (افتراضي 1)
            material: نوع خامة المنتج (مثلاً: قطن 100%، ميلتون، كتان، جبردين، وتربروف) إن وجدت
        """
        if price <= 0:
            return {"status": "error", "reason": "السعر لازم يكون أكبر من صفر"}
        if qty <= 0:
            return {"status": "error", "reason": "الكمية لازم تكون أكبر من صفر"}

        proposals.append({
            "type": "new_product",
            "name": name.strip(),
            "price": float(price),
            "size": size.strip().upper(),
            "color": color.strip(),
            "qty": int(qty),
            "material": material.strip(),
        })
        return {"status": "waiting_confirmation", "product": name, "price": price}

    def propose_new_branch(name: str, governorate: str, address: str, google_maps_url: str = "", working_hours: str = "من 12 ظهراً لـ 11 مساءً") -> dict:
        """بتجهّز إضافة فرع أو عنوان للمحل وتستنى تأكيد صاحب المحل بالأزرار.

        Args:
            name: اسم الفرع (مثلاً: فرع مدينة نصر، فرع الدقي)
            governorate: المحافظة (القاهرة، الجيزة، الإسكندرية...)
            address: العنوان بالتفصيل
            google_maps_url: رابط اللوكيشن على خرائط جوجل (إن وجد)
            working_hours: مواعيد العمل
        """
        proposals.append({
            "type": "new_branch",
            "name": name.strip(),
            "governorate": governorate.strip(),
            "address": address.strip(),
            "google_maps_url": google_maps_url.strip(),
            "working_hours": working_hours.strip(),
        })
        return {"status": "waiting_confirmation", "branch": name}

    def propose_new_promo(code: str = "", promo_type: str = "percentage", discount_value: float = 0.0, description: str = "", min_order_items: int = 1, min_order_value: float = 0.0) -> dict:
        """بتجهّز إضافة كود خصم أو عرض ترويجي أو كارت هدايا/فاوتشر للزيارة القادمة وتستنى تأكيد التاجر بالأزرار.

        Args:
            code: كود الخصم (زي EID20 أو GIFT200) أو اتركه فارغاً لو عرض عام تلقائي (زي اشتري بـ 1000ج واحصل على فاوتشر 200ج)
            promo_type: نوع العرض (voucher = كارت هدايا/فاوتشر، percentage = نسبة مئوية، fixed = خصم بالجنيه، free_shipping = شحن مجاني)
            discount_value: قيمة الخصم أو الفاوتشر (مثلاً 200 لو فاوتشر 200 جنيه أو 20 لو 20%)
            description: وصف العرض بالعربي (مثلاً: اشتري بـ 1000 واحصل على فاوتشر 200 لطلبك القادم)
            min_order_items: الحد الأدنى لعدد القطع لتطبيق العرض (افتراضي 1)
            min_order_value: الحد الأدنى لقيمة الأوردر لتطبيق العرض (مثلاً 1000 لعروض الفاوتشر)
        """
        proposals.append({
            "type": "new_promo",
            "code": code.strip().upper(),
            "promo_type": promo_type.strip(),
            "discount_value": float(discount_value),
            "description": description.strip() or f"عرض {code}",
            "min_order_items": int(min_order_items),
            "min_order_value": float(min_order_value),
        })
        return {"status": "waiting_confirmation", "promo": description}

    def propose_cancel_order(order_number: int, reason: str = "") -> dict:
        """بتجهّز إلغاء أوردر وإرجاع بضاعته للمخزن وتستنى تأكيد صاحب المحل بالأزرار. مابتنفذش حاجة مباشرة.

        Args:
            order_number: رقم الأوردر المطلوب إلغاؤه (مثلاً 5)
            reason: سبب الإلغاء (اختياري)
        """
        proposals.append({
            "type": "cancel_order",
            "order_number": int(order_number),
            "reason": reason.strip(),
        })
        return {"status": "waiting_confirmation", "order_number": order_number}

    def propose_update_order(order_number: int, phone: str = "", address: str = "", status: str = "") -> dict:
        """بتجهّز تعديل بيانات أوردر (العنوان، الموبايل، أو الحالة) وتستنى تأكيد صاحب المحل بالأزرار. مابتنفذش حاجة مباشرة.

        Args:
            order_number: رقم الأوردر المراد تعديله
            phone: رقم الموبايل الجديد (إن وجد)
            address: العنوان الجديد (إن وجد)
            status: الحالة الجديدة (confirmed, shipped, delivered) إن وجدت
        """
        clean_status = status.strip().lower()
        if clean_status and clean_status not in ("confirmed", "shipped", "delivered"):
            return {
                "status": "error",
                "reason": "الحالة غير صالحة. الحالات المسموح بها هي: confirmed, shipped, delivered"
            }

        proposals.append({
            "type": "update_order",
            "order_number": int(order_number),
            "phone": phone.strip(),
            "address": address.strip(),
            "status": clean_status,
        })
        return {"status": "waiting_confirmation", "order_number": order_number}

    def propose_set_product_image(product_name: str) -> dict:
        """بتجهّز ربط آخر صورة بعتها صاحب المحل بمنتج موجود بالفعل (بتستبدل صورته) وتستنى تأكيده بالأزرار.

        Args:
            product_name: اسم المنتج الموجود في المخزن (مثلاً: هودي، ترنج شتوي)
        """
        proposals.append({"type": "set_image", "product_name": product_name.strip()})
        return {"status": "waiting_confirmation", "product": product_name}

    tools = [
        check_stock,
        list_inventory,
        propose_stock_change,
        propose_new_product,
        propose_new_branch,
        propose_new_promo,
        propose_cancel_order,
        propose_update_order,
        propose_set_product_image,
    ]
    return tools


def build_proposals_summary(proposals: list) -> str:
    """بناء نص ملخص العمليات المعروضة في رسالة طلب التأكيد"""
    kinds = {"sale": "خصم (بيع من المحل)", "restock": "إضافة بضاعة", "return": "مرتجع"}
    summary_lines = []
    for p in proposals:
        ptype = p.get("type")
        if ptype == "new_product":
            mat_txt = f" | خامة: {p['material']}" if p.get("material") else ""
            summary_lines.append(f"✨ إضافة صنف: {p['name']} ({p['color']} - {p['size']}{mat_txt}) | السعر: {p['price']:.0f}ج | الكمية: {p['qty']}")
        elif ptype == "new_branch":
            summary_lines.append(f"📍 إضافة فرع: {p['name']} ({p['governorate']} - {p['address']})")
        elif ptype == "new_promo":
            code_txt = f" (كود: {p['code']})" if p.get('code') else ""
            summary_lines.append(f"🎁 إضافة عرض/خصم: {p['description']}{code_txt}")
        elif ptype == "set_image":
            summary_lines.append(f"🖼️ تغيير صورة المنتج: {p['product_name']}")
        elif ptype == "cancel_order":
            reason_txt = f" (السبب: {p['reason']})" if p.get("reason") else ""
            summary_lines.append(f"❌ إلغاء الأوردر #{p['order_number']}{reason_txt} وإرجاع بضاعته للمخزن")
        elif ptype == "update_order":
            ch = []
            if p.get("phone"):
                ch.append(f"موبايل: {p['phone']}")
            if p.get("address"):
                ch.append(f"عنوان: {p['address']}")
            if p.get("status"):
                ch.append(f"حالة: {p['status']}")
            ch_str = " | ".join(ch) if ch else "تحديث بيانات"
            summary_lines.append(f"📝 تعديل الأوردر #{p['order_number']} ({ch_str})")
        else:
            summary_lines.append(f"• {kinds.get(p.get('kind'), 'تعديل')}: {abs(p.get('change', 0))} × {p.get('label', '')}")
    return "\n".join(summary_lines)


async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    tenant = await asyncio.to_thread(get_owner_tenant, chat_id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    tenant_id = tenant["id"]
    proposals = []
    photo = get_fresh_photo(chat_id)
    tools = build_tools(tenant_id, proposals, has_image=bool(photo))

    reply = await asyncio.to_thread(run_llm, chat_id, update.message.text, tools)
    await reply_with_proposals(update, context, chat_id, tenant_id, proposals, reply, image=photo)


async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """استقبال صور المنتجات (صورة عادية أو صورة مبعوتة كملف) مع الكابشن كتعليمات."""
    msg = update.message
    chat_id = update.effective_chat.id
    tenant = await asyncio.to_thread(get_owner_tenant, chat_id)
    if not tenant:
        await msg.reply_text(NOT_OWNER_MESSAGE)
        return
    tenant_id = tenant["id"]

    if msg.photo:
        tg_obj = msg.photo[-1]  # أكبر نسخة
        mime = "image/jpeg"
    elif msg.document:
        tg_obj = msg.document
        fname = (msg.document.file_name or "").lower()
        if fname.endswith(".png"):
            mime = "image/png"
        elif fname.endswith(".webp"):
            mime = "image/webp"
        elif fname.endswith(".jpg") or fname.endswith(".jpeg"):
            mime = "image/jpeg"
        else:
            mime = msg.document.mime_type or "image/jpeg"
    else:
        return

    if mime not in ALLOWED_IMAGE_MIMES:
        await msg.reply_text("❌ الصيغة دي مش مدعومة. ابعت الصورة بصيغة JPG أو PNG أو WEBP.")
        return
    if getattr(tg_obj, "file_size", 0) and tg_obj.file_size > MAX_IMAGE_BYTES:
        await msg.reply_text("❌ الصورة كبيرة (الحد الأقصى 5 ميجا). ابعتها كصورة عادية مش كملف.")
        return

    status_msg = await msg.reply_text("🖼️ جاري تحليل الصورة...")
    try:
        tg_file = await context.bot.get_file(tg_obj.file_id)
        data = bytes(await tg_file.download_as_bytearray())
    except Exception:
        await status_msg.edit_text("❌ معرفتش أحمّل الصورة، ابعتها تاني.")
        return

    photo = {"bytes": data, "mime": mime, "ts": time.time()}
    remember_photo(chat_id, data, mime)

    # فحص لو التاجر اختار مسبقاً صنف لتحديث صورته بأمر /set_image
    waiting_prod = context.user_data.pop("waiting_image_product", None) if context.user_data else None
    if waiting_prod:
        proposals = [{"type": "set_image", "product_name": waiting_prod["name"]}]
        try:
            await status_msg.delete()
        except Exception:
            pass
        await reply_with_proposals(
            update, context, chat_id, tenant_id, proposals,
            reply=f"استلمت الصورة، وهربطها بمنتج «{waiting_prod['name']}».",
            intro=f"🖼️ استلمت الصورة الجديدة لمنتج «{waiting_prod['name']}»:",
            image=photo,
        )
        return

    caption = (msg.caption or "").strip()
    proposals = []
    tools = build_tools(tenant_id, proposals, has_image=True)
    prompt = (
        f"صاحب المحل بعتلك صورة منتج. كلامه معاها: «{caption}». نفّذ المطلوب واستدعِ الأدوات المناسبة."
        if caption else
        "صاحب المحل بعتلك صورة منتج من غير أي كلام. بص عليها وقوله شايف إيه، واسأله عن الناقص (الاسم، السعر، المقاس، الكمية)."
    )
    content = [types.Part.from_bytes(data=data, mime_type=mime), prompt]

    reply = await asyncio.to_thread(run_llm, chat_id, content, tools)
    try:
        await status_msg.delete()
    except Exception:
        pass
    await reply_with_proposals(
        update, context, chat_id, tenant_id, proposals, reply,
        intro="استلمت الصورة، وهنفّذ العملية دي:", image=photo,
    )


async def handle_voice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """استقبال وتحليل الفويسات الصوتية مع حد 60 ثانية واستخدام mime_type الديناميكي"""
    chat_id = update.effective_chat.id
    tenant = get_owner_tenant(chat_id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    tenant_id = tenant["id"]
    voice = update.message.voice or update.message.audio
    if not voice:
        return

    # حد أقصى 60 ثانية للفويس
    MAX_VOICE_DURATION_SECONDS = 60
    if getattr(voice, "duration", 0) and voice.duration > MAX_VOICE_DURATION_SECONDS:
        await update.message.reply_text(
            f"⏱️ مدة التسجيل الصوتي ({voice.duration} ثانية) تتجاوز الحد الأقصى المسموح به (60 ثانية). يرجى إرسال تسجيل أقصر."
        )
        return

    status_msg = await update.message.reply_text("🎙️ جاري الاستماع للفويس وتحليله...")

    voice_file = await context.bot.get_file(voice.file_id)
    audio_bytes = await voice_file.download_as_bytearray()

    proposals = []
    photo = get_fresh_photo(chat_id)
    tools = build_tools(tenant_id, proposals, has_image=bool(photo))

    mime_type = getattr(voice, "mime_type", None) or "audio/ogg"
    audio_part = types.Part.from_bytes(data=bytes(audio_bytes), mime_type=mime_type)
    content_list = [audio_part, "استمع للريكورد ده من صاحب المحل ونفّذ المطلوب واستدعِ الأدوات المناسبة."]

    reply = await asyncio.to_thread(run_llm, chat_id, content_list, tools)
    await status_msg.delete()
    await reply_with_proposals(
        update, context, chat_id, tenant_id, proposals, reply,
        intro="سمعت الفويس، وهنفّذ العملية دي:", image=photo,
    )


def process_confirmed_proposals(tenant_id: str, pid: str, items: list, image: dict = None) -> dict:
    """
    تنفيذ العمليات بعد تأكيد صاحب المحل:
    - البحث عن الصنف بـ (product_id, size, _norm_ar(color)).
    - لو موجود: تزويد كميته بـ apply_stock_movement (restock).
    - لو غير موجود: إنشاؤه بـ stock=0 وكود لاتيني P[A-Z]{6} ثم apply_stock_movement بالكمية.
    - لو السعر مختلف: تنبيه التاجر مع الإبقاء على السعر الأساسي للمنتج.
    - تنفيذ cancel_order و update_order.
    - إرجاع تقرير رقمي مفصل بالأعداد الناجحة والفاشلة.
    """
    successes = []
    failures = []
    warnings = []

    # رفع الصورة مرة واحدة لو في عمليات محتاجاها
    image_url = None
    if image and any(it.get("type") in ("new_product", "set_image") for it in items):
        try:
            image_url = upload_product_image(tenant_id, image["bytes"], image["mime"])
        except Exception as e:
            print("فشل رفع الصورة:", e)
            warnings.append(f"الصورة ماترفعتش ({e}). العمليات اتنفذت من غير صورة.")

    for i, item in enumerate(items):
        itype = item.get("type")
        ref_code = f"tg-{pid}-{i}"

        if itype == "new_promo":
            try:
                p_payload = {
                    "tenant_id": tenant_id,
                    "code": item.get("code") or None,
                    "promo_type": item["promo_type"],
                    "discount_value": item.get("discount_value", 0),
                    "min_order_items": item.get("min_order_items", 1),
                    "description": item["description"],
                    "is_active": True,
                }
                if item.get("min_order_value") and float(item["min_order_value"]) > 0:
                    p_payload["min_order_value"] = float(item["min_order_value"])

                try:
                    supabase.table("promotions").insert(p_payload).execute()
                except Exception as e_col:
                    err_s = str(e_col)
                    if "promotions_type_check" in err_s:
                        p_payload["promo_type"] = "fixed"
                        if "[كارت هدايا]" not in p_payload["description"] and "كارت" not in p_payload["description"]:
                            p_payload["description"] = f"[كارت هدايا] {p_payload['description']}"
                        try:
                            supabase.table("promotions").insert(p_payload).execute()
                        except Exception as e_retry:
                            err_s = str(e_retry)

                    if "min_order_value" in err_s:
                        p_payload.pop("min_order_value", None)
                        supabase.table("promotions").insert(p_payload).execute()
                    else:
                        raise e_col

                successes.append(f"عرض ترويجي: {item['description']}")
            except Exception as e:
                print("فشل إضافة العرض:", e)
                failures.append(f"عرض ترويجي: {item['description']} ({e})")

        elif itype == "new_branch":
            try:
                supabase.table("branches").insert({
                    "tenant_id": tenant_id,
                    "name": item["name"],
                    "governorate": item["governorate"],
                    "address": item["address"],
                    "google_maps_url": item.get("google_maps_url"),
                    "working_hours": item.get("working_hours", "من 12 ظهراً لـ 11 مساءً"),
                }).execute()
                successes.append(f"فرع جديد: {item['name']} ({item['address']})")
            except Exception as e:
                print("فشل إضافة الفرع:", e)
                failures.append(f"فرع: {item['name']} ({e})")

        elif itype == "set_image":
            try:
                if not image_url:
                    raise RuntimeError("مفيش صورة جاهزة للربط")
                prod = _find_product_by_name(tenant_id, item["product_name"])
                if not prod:
                    raise RuntimeError("المنتج مش موجود في المخزن")
                supabase.table("products").update({"image_url": image_url}).eq("id", prod["id"]).execute()
                supabase.table("variants").update({"image_url": image_url}).eq("tenant_id", tenant_id).eq("product_id", prod["id"]).execute()
                successes.append(f"تم تحديث صورة «{prod['name']}» 🖼️")
            except Exception as e:
                failures.append(f"صورة «{item['product_name']}»: {e}")

        elif itype == "cancel_order":
            try:
                res = cancel_order(tenant_id, item["order_number"], reason=item.get("reason", ""))
                if res.get("status") == "ok":
                    successes.append(f"إلغاء الأوردر #{item['order_number']}")
                else:
                    failures.append(f"إلغاء الأوردر #{item['order_number']}: {res.get('reason', 'فشل الإلغاء')}")
            except Exception as e:
                failures.append(f"إلغاء الأوردر #{item['order_number']}: {e}")

        elif itype == "update_order":
            up_errs = []
            if item.get("phone") or item.get("address"):
                try:
                    res_sh = update_order_shipping(
                        tenant_id, item["order_number"], phone=item.get("phone", ""), address=item.get("address", "")
                    )
                    if res_sh.get("status") != "ok":
                        up_errs.append(res_sh.get("reason", "فشل تحديث الشحن"))
                except Exception as e:
                    up_errs.append(str(e))
            if item.get("status"):
                try:
                    res_st = update_order_status(tenant_id, item["order_number"], new_status=item["status"])
                    if res_st.get("status") != "ok":
                        up_errs.append(res_st.get("reason", "فشل تحديث الحالة"))
                except Exception as e:
                    up_errs.append(str(e))

            if up_errs:
                failures.append(f"تعديل الأوردر #{item['order_number']}: {', '.join(up_errs)}")
            else:
                successes.append(f"تعديل بيانات الأوردر #{item['order_number']}")

        elif itype == "new_product":
            try:
                # 1. فحص المنتج بالاسم ومقارنته بـ _norm_ar
                prod_rows = (
                    supabase.table("products")
                    .select("id, name, price, image_url, material")
                    .eq("tenant_id", tenant_id)
                    .execute()
                    .data or []
                )
                target_prod_norm = _norm_ar(item["name"])
                matched_prod = None
                for pr in prod_rows:
                    if _norm_ar(pr.get("name", "")) == target_prod_norm:
                        matched_prod = pr
                        break

                if matched_prod:
                    product_id = matched_prod["id"]
                    prod_updates = {}
                    if image_url and not matched_prod.get("image_url"):
                        prod_updates["image_url"] = image_url
                    if item.get("material") and not matched_prod.get("material"):
                        prod_updates["material"] = item["material"]
                    if prod_updates:
                        supabase.table("products").update(prod_updates).eq("id", product_id).execute()

                    existing_price = float(matched_prod.get("price") or 0)
                    if abs(existing_price - float(item["price"])) > 0.01:
                        warnings.append(
                            f"تنبيه: منتج «{item['name']}» مسجل مسبقاً بسعر {existing_price:.0f}ج وسعرك المدخل {item['price']:.0f}ج. تم الإبقاء على السعر الأصلي للمنتج."
                        )
                else:
                    new_prod_payload = {
                        "tenant_id": tenant_id,
                        "name": item["name"],
                        "price": item["price"],
                        "active": True,
                        **({"image_url": image_url} if image_url else {}),
                    }
                    if item.get("material"):
                        new_prod_payload["material"] = item["material"]

                    new_p = (
                        supabase.table("products")
                        .insert(new_prod_payload)
                        .execute()
                        .data
                    )
                    if not new_p:
                        raise RuntimeError(f"فشل إنشاء المنتج {item['name']}")
                    product_id = new_p[0]["id"]

                # 2. فحص هل الصنف variant موجود بنفس (product_id, size, _norm_ar(color))
                var_rows = (
                    supabase.table("variants")
                    .select("id, sku, size, color, stock")
                    .eq("tenant_id", tenant_id)
                    .eq("product_id", product_id)
                    .execute()
                    .data or []
                )
                target_size = item["size"].strip().upper()
                target_color_norm = _norm_ar(item["color"])
                matched_var = None
                for vr in var_rows:
                    if vr.get("size", "").strip().upper() == target_size and _norm_ar(vr.get("color", "")) == target_color_norm:
                        matched_var = vr
                        break

                if matched_var:
                    # صنف موجود بالفعل -> زيادة رصيده بحركة restock
                    var_id = matched_var["id"]
                    supabase.rpc(
                        "apply_stock_movement",
                        {
                            "p_tenant": tenant_id,
                            "p_variant": var_id,
                            "p_change": item["qty"],
                            "p_reason": "restock",
                            "p_source": "telegram",
                            "p_ref": ref_code,
                        },
                    ).execute()
                    if image_url:
                        supabase.table("variants").update({"image_url": image_url}).eq("id", var_id).execute()
                    successes.append(f"تزويد مخزون «{item['name']}» ({item['color']} {item['size']}): +{item['qty']}")
                else:
                    # صنف غير موجود -> إنشاؤه بـ stock=0 وكود لاتيني، ثم تزويد رصيده بحركة restock
                    sku = generate_sku()
                    new_v = (
                        supabase.table("variants")
                        .insert({
                            "tenant_id": tenant_id,
                            "product_id": product_id,
                            "sku": sku,
                            "size": target_size,
                            "color": item["color"].strip(),
                            "stock": 0,
                            "reserved": 0,
                            "active": True,
                            **({"image_url": image_url} if image_url else {}),
                        })
                        .execute()
                        .data
                    )
                    if not new_v:
                        raise RuntimeError(f"فشل إنشاء الصنف {item['name']} {sku}")
                    var_id = new_v[0]["id"]

                    supabase.rpc(
                        "apply_stock_movement",
                        {
                            "p_tenant": tenant_id,
                            "p_variant": var_id,
                            "p_change": item["qty"],
                            "p_reason": "restock",
                            "p_source": "telegram",
                            "p_ref": ref_code,
                        },
                    ).execute()
                    successes.append(f"صنف جديد «{item['name']}» ({item['color']} {item['size']}) بكود {sku}: كمية {item['qty']}")

            except Exception as e:
                print("فشل إضافة/تحديث الصنف:", e)
                failures.append(f"صنف «{item.get('name', '')}» ({e})")

        else:
            # حركة تعديل مخزون عادية (sale / restock / return)
            try:
                supabase.rpc(
                    "apply_stock_movement",
                    {
                        "p_tenant": tenant_id,
                        "p_variant": item["variant_id"],
                        "p_change": item["change"],
                        "p_reason": item["kind"],
                        "p_source": "telegram",
                        "p_ref": ref_code,
                    },
                ).execute()
                successes.append(item.get("label", "تعديل مخزون"))
            except Exception as e:
                print("فشل تعديل المخزون:", type(e).__name__)
                failures.append(f"{item.get('label', 'صنف')}: {e}")

    return {
        "success_count": len(successes),
        "failure_count": len(failures),
        "successes": successes,
        "failures": failures,
        "warnings": warnings,
    }


async def set_image_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """أمر مخصص للتاجر لتغيير أو إضافة صورة لأي منتج مسجل في المحل أو استبدال صورة قديمة"""
    tenant = get_owner_tenant(update.effective_chat.id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    tenant_id = tenant["id"]
    chat_id = update.effective_chat.id
    target_name = " ".join(context.args).strip() if context.args else ""

    if not target_name:
        prods = (
            supabase.table("products")
            .select("id, name, image_url")
            .eq("tenant_id", tenant_id)
            .order("name", desc=False)
            .limit(20)
            .execute()
            .data or []
        )
        if not prods:
            await update.message.reply_text("📦 مفيش منتجات مسجلة في المخزن حالياً. ضيف منتجات الأول!")
            return

        keyboard = []
        row = []
        for p in prods:
            has_img = "🖼️" if p.get("image_url") else "📷"
            btn_text = f"{has_img} {p['name'][:18]}"
            row.append(InlineKeyboardButton(btn_text, callback_data=f"chgimg:{p['id'][:10]}"))
            if len(row) == 2:
                keyboard.append(row)
                row = []
        if row:
            keyboard.append(row)

        await update.message.reply_text(
            "📸 *تغيير أو إضافة صورة لمنتج:*\n\n"
            "اضغط على الصنف اللي عايز تغير صورته أو ابعت صورة واكتب في الرسالة «غير صورة اسم_المنتج»:\n"
            "• 🖼️ = له صورة حالية\n"
            "• 📷 = بدون صورة حالياً",
            reply_markup=InlineKeyboardMarkup(keyboard),
            parse_mode="Markdown"
        )
        return

    prod = _find_product_by_name(tenant_id, target_name)
    if not prod:
        await update.message.reply_text(f"❌ لم يتم العثور على منتج باسم «{target_name}» في المخزن. تأكد من الاسم أو اكتب /stock لمعرفة الأسماء.")
        return

    photo = get_fresh_photo(chat_id)
    if photo:
        proposals = [{"type": "set_image", "product_name": prod["name"]}]
        await reply_with_proposals(
            update, context, chat_id, tenant_id, proposals,
            reply=f"استلمت الصورة، وهربطها بمنتج «{prod['name']}».",
            intro="جاهز لتحديث صورة المنتج:",
            image=photo
        )
    else:
        if context.user_data is not None:
            context.user_data["waiting_image_product"] = {"id": prod["id"], "name": prod["name"]}
        buttons = []
        if prod.get("image_url"):
            buttons.append([InlineKeyboardButton("🗑️ حذف الصورة الحالية", callback_data=f"delimg:{prod['id'][:10]}")])
        reply_markup = InlineKeyboardMarkup(buttons) if buttons else None
        await update.message.reply_text(
            f"📸 *جاهز لتحديث صورة «{prod['name']}»!*\n\n"
            "ابعت الصورة دلوقتي (صورة عادية أو ملف)، وهحدثهالك فوراً في المخزن والرد التلقائي 🚀\n"
            "(لو غيرت رأيك أو بعت صورة غلط تقدر تغيرها أو تحذفها في أي وقت).",
            reply_markup=reply_markup,
            parse_mode="Markdown"
        )


async def on_button(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    action, _, pid = query.data.partition(":")
    chat_id = query.message.chat.id
    tenant = get_owner_tenant(chat_id)

    # معالجة اختيار تغيير صورة المنتج من قائمة /set_image
    if action == "chgimg":
        short_id = pid
        if not tenant:
            await query.answer("غير مصرح", show_alert=True)
            return
        prods = (
            supabase.table("products")
            .select("id, name, image_url")
            .eq("tenant_id", tenant["id"])
            .execute()
            .data or []
        )
        matched = next((p for p in prods if str(p.get("id", "")).startswith(short_id)), None)
        if not matched:
            await query.answer("لم يتم العثور على المنتج", show_alert=True)
            return

        photo = get_fresh_photo(chat_id)
        if photo:
            proposals = [{"type": "set_image", "product_name": matched["name"]}]
            await reply_with_proposals(
                update, context, chat_id, tenant["id"], proposals,
                reply=f"استلمت الصورة، وهربطها بمنتج «{matched['name']}».",
                intro="جاهز لتحديث صورة المنتج:",
                image=photo
            )
        else:
            if context.user_data is not None:
                context.user_data["waiting_image_product"] = {"id": matched["id"], "name": matched["name"]}
            buttons = []
            if matched.get("image_url"):
                buttons.append([InlineKeyboardButton("🗑️ حذف الصورة الحالية", callback_data=f"delimg:{matched['id'][:10]}")])
            reply_markup = InlineKeyboardMarkup(buttons) if buttons else None
            await query.edit_message_text(
                f"📸 تم اختيار «{matched['name']}»!\n\n"
                "ابعت الصورة الجديدة دلوقتي كصورة أو ملف، وهحدثهالك فوراً في المخزن والرد التلقائي 🚀\n"
                "(لو بعت صورة بالغلط تقدر تغيرها في أي وقت بنفس الطريقة).",
                reply_markup=reply_markup
            )
        return

    # معالجة حذف صورة المنتج
    if action == "delimg":
        short_id = pid
        if not tenant:
            await query.answer("غير مصرح", show_alert=True)
            return
        prods = (
            supabase.table("products")
            .select("id, name, image_url")
            .eq("tenant_id", tenant["id"])
            .execute()
            .data or []
        )
        matched = next((p for p in prods if str(p.get("id", "")).startswith(short_id)), None)
        if not matched:
            await query.answer("لم يتم العثور على المنتج", show_alert=True)
            return

        try:
            supabase.table("products").update({"image_url": None}).eq("id", matched["id"]).eq("tenant_id", tenant["id"]).execute()
            supabase.table("variants").update({"image_url": None}).eq("product_id", matched["id"]).eq("tenant_id", tenant["id"]).execute()
            await query.edit_message_text(f"🗑️ تم حذف صورة «{matched['name']}» بنجاح وإزالتها من الرد التلقائي.")
        except Exception as e:
            await query.edit_message_text(f"❌ تعذر حذف الصورة: {e}")
        return

    pending = pop_pending_operation(context, pid)
    if not pending or pending["chat_id"] != chat_id or not tenant:
        await query.edit_message_text("العملية دي انتهت صلاحيتها أو اتنفذت قبل كده.")
        return

    if action == "no":
        await query.edit_message_text("اتلغت العملية ❌")
        return

    tenant_id = pending.get("tenant_id") or tenant["id"]

    # تنفيذ الاستدعاءات المتزامنة داخل thread لعدم حظر الـ Event Loop
    extra = (pending["image"],) if pending.get("image") else ()
    report = await asyncio.to_thread(
        process_confirmed_proposals, tenant_id, pid, pending["items"], *extra
    )

    # طباعة التفاصيل الكاملة للأخطاء في السجلات/الـ console فقط
    if report.get("failures"):
        print(f"⚠️ [on_button failures] pid={pid} tenant={tenant_id}: {report['failures']}")

    lines = ["📊 تقرير تنفيذ العمليات:"]
    lines.append(f"✅ الناجح: {report['success_count']} عملية")
    lines.append(f"❌ الفاشل: {report['failure_count']} عملية\n")

    if report["successes"]:
        lines.append("📌 تم بنجاح:")
        for s in report["successes"][:12]:
            lines.append(f"• {s}")
        if len(report["successes"]) > 12:
            lines.append(f"... والمزيد ({len(report['successes']) - 12} عملية أخرى)")

    if report["failure_count"] > 0:
        lines.append("\n⚠️ تعذر تنفيذ بعض العمليات، يرجى مراجعة البيانات والمحاولة لاحقاً.")

    if report.get("warnings"):
        lines.append("\n🔔 تنبيهات:")
        for w in report["warnings"][:5]:
            lines.append(f"• {w}")

    # التقرير بنص عادي بدون parse_mode لتفادي أخطاء التنسيق
    await query.edit_message_text("\n".join(lines))


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """
    استقبال ملف Excel أو CSV من التاجر:
    - الحد الأقصى للحجم: 2 ميجابايت
    - الحد الأقصى للصفوف: 100 صف في الدفعة الواحدة
    - الامتدادات المقبولة: .xlsx و .csv فقط (تم حذف .xls)
    """
    chat_id = update.effective_chat.id
    tenant = get_owner_tenant(chat_id)
    if not tenant:
        await update.message.reply_text(NOT_OWNER_MESSAGE)
        return

    doc = update.message.document
    if not doc:
        return

    fname = (doc.file_name or "").lower()
    # لو التاجر بعت صورة كملف (Document) بدل صورة عادية: حوّلها لمعالج الصور فوراً
    if (doc.mime_type or "").startswith("image/") or any(fname.endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp"]):
        return await handle_photo(update, context)

    # فحص حجم الملف (حد 2 ميجا)
    MAX_DOC_SIZE_BYTES = 2 * 1024 * 1024
    if doc.file_size and doc.file_size > MAX_DOC_SIZE_BYTES:
        await update.message.reply_text("❌ حجم الملف كبير جداً! الحد الأقصى المسموح به هو 2 ميجابايت.")
        return

    if not (fname.endswith(".xlsx") or fname.endswith(".csv")):
        await update.message.reply_text(
            "📋 ابعتلي ملف Excel بصيغة (.xlsx) أو CSV بصيغة (.csv) فقط (حتى 2 ميجا و100 صف في الدفعة).\n"
            "الأعمدة المطلوبة:\n"
            "اسم المنتج | السعر | المقاس | اللون | الكمية\n\n"
            "💡 ولو عايز تضيف منتج بصورته: ابعت الصورة وفي الكابشن اكتب مثلاً «وصل 50 ترنج بني مقاس L بـ600»."
        )
        return

    status_msg = await update.message.reply_text("📊 جاري قراءة الشيت...")

    file = await context.bot.get_file(doc.file_id)
    file_bytes = await file.download_as_bytearray()

    rows_parsed = []
    errors = []

    try:
        if fname.endswith(".csv"):
            import csv
            text = file_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text))
            raw_rows = list(reader)
        else:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(bytes(file_bytes)), data_only=True, read_only=True)
            ws = wb.active
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not header_row:
                await status_msg.edit_text("❌ الملف فاضي!")
                return
            headers = [str(c or "").strip() for c in header_row]
            raw_rows = []
            for row in ws.iter_rows(min_row=2, max_row=502, values_only=True):
                raw_rows.append(dict(zip(headers, [str(v or "").strip() for v in row])))
    except Exception as e:
        await status_msg.edit_text(f"❌ فشل قراءة الملف: {type(e).__name__}")
        return

    # حد 100 صف في الدفعة الواحدة (بعد شيل الصفوف الفاضية)
    MAX_DOC_ROWS = 100
    raw_rows = [r for r in raw_rows if any(str(v or "").strip() for v in r.values())]
    if len(raw_rows) > MAX_DOC_ROWS:
        await status_msg.edit_text(
            f"❌ الملف يحتوي على {len(raw_rows)} صف! الحد الأقصى المسموح به في الدفعة الواحدة هو {MAX_DOC_ROWS} صف."
        )
        return

    # شيل الصفوف الفاضية (إكسيل بيرجّع صفوف فاضية لحد max_row وكانت بتتحسب ضمن حد الـ100)
    raw_rows = [r for r in raw_rows if any(str(v or "").strip() for v in r.values())]

    COL_MAP = {
        "name":  ["اسم المنتج", "المنتج", "الاسم", "الصنف", "name", "product"],
        "price": ["السعر", "سعر البيع", "price"],
        "size":  ["المقاس", "مقاس", "size"],
        "color": ["اللون", "لون", "color"],
        "qty":   ["الكمية", "كمية", "عدد", "qty", "quantity", "stock"],
        "material": ["الخامة", "خامة", "نوع الخامة", "القماش", "material", "fabric"],
    }

    def find_col(row_keys, aliases):
        wanted = {a.strip().lower() for a in aliases}
        for k in row_keys:
            if (k or "").strip().lower() in wanted:
                return k
        return None

    if not raw_rows:
        await status_msg.edit_text("❌ الملف فاضي!")
        return

    first_keys = list(raw_rows[0].keys())
    col_name  = find_col(first_keys, COL_MAP["name"])
    col_price = find_col(first_keys, COL_MAP["price"])
    col_size  = find_col(first_keys, COL_MAP["size"])
    col_color = find_col(first_keys, COL_MAP["color"])
    col_qty   = find_col(first_keys, COL_MAP["qty"])
    col_material = find_col(first_keys, COL_MAP["material"])

    missing = [k for k, v in {"اسم المنتج": col_name, "السعر": col_price, "المقاس": col_size, "اللون": col_color}.items() if not v]
    if missing:
        await status_msg.edit_text(
            f"❌ الأعمدة دي مش موجودة في الشيت: {', '.join(missing)}\n"
            f"الأعمدة اللي عندك: {', '.join(first_keys)}\n\n"
            "الأعمدة المطلوبة:\n"
            "اسم المنتج | السعر | المقاس | اللون | الكمية | الخامة (اختياري)"
        )
        return

    for i, row in enumerate(raw_rows, start=2):
        name  = (row.get(col_name) or "").strip()
        price_raw = (row.get(col_price) or "0").strip().replace(",", "")
        size  = (row.get(col_size) or "").strip().upper()
        color = (row.get(col_color) or "").strip()
        qty_raw = (row.get(col_qty) or "1").strip() if col_qty else "1"
        material = (row.get(col_material) or "").strip() if col_material else ""

        if not name or not size or not color:
            errors.append(f"سطر {i}: بيانات ناقصة — تم تخطيه")
            continue

        try:
            price = float(price_raw or 0)
            qty   = int(float(qty_raw or 1))
        except ValueError:
            errors.append(f"سطر {i} ({name}): سعر أو كمية غير صالحة")
            continue

        if price <= 0:
            errors.append(f"سطر {i} ({name}): السعر لازم أكبر من صفر")
            continue

        rows_parsed.append({
            "type":  "new_product",
            "name":  name,
            "price": price,
            "size":  size,
            "color": color,
            "qty":   max(1, qty),
            "material": material,
        })

    if not rows_parsed:
        err_text = "\n".join(errors[:5])
        await status_msg.edit_text(f"❌ مفيش صفوف صالحة في الشيت!\n{err_text}")
        return

    pid = uuid.uuid4().hex[:12]
    store_pending_operation(context, pid, {
        "chat_id": chat_id,
        "tenant_id": tenant["id"],
        "items": rows_parsed,
    })

    preview_lines = [f"📊 لقيت {len(rows_parsed)} صنف في الشيت:\n"]
    for p in rows_parsed[:12]:
        preview_lines.append(f"• {p['name']} | {p['color']} {p['size']} | {p['price']:.0f}ج | كمية: {p['qty']}")
    if len(rows_parsed) > 12:
        preview_lines.append(f"... و{len(rows_parsed) - 12} صنف تاني")

    if errors:
        preview_lines.append(f"\n⚠️ {len(errors)} سطر تم تخطيه:")
        for e in errors[:3]:
            preview_lines.append(f"  - {e}")

    preview_lines.append("\nتأكيد إضافة كل ده للمخزن؟")

    keyboard = InlineKeyboardMarkup([[
        InlineKeyboardButton("✅ تأكيد الكل", callback_data=f"ok:{pid}"),
        InlineKeyboardButton("❌ إلغاء", callback_data=f"no:{pid}"),
    ]])

    await status_msg.edit_text("\n".join(preview_lines), reply_markup=keyboard)


async def telegram_error_handler(update: object, context: ContextTypes.DEFAULT_TYPE) -> None:
    """معالج عام للأخطاء يسجل نوع الخطأ ويحمي التوكنات والبيانات الحساسة"""
    err = context.error
    err_type = type(err).__name__
    err_msg = str(err)
    if BOT_TOKEN and BOT_TOKEN in err_msg:
        err_msg = err_msg.replace(BOT_TOKEN, "[PROTECTED_TOKEN]")
    logging.error("Telegram Bot Error: %s: %s", err_type, err_msg)


def main():
    asyncio.set_event_loop(asyncio.new_event_loop())

    async def post_init(application):
        async def background_checker():
            while True:
                await asyncio.sleep(1800)  # فحص كل 30 دقيقة
                try:
                    await check_subscriptions_and_notify(application)
                except Exception as err:
                    print("Error in subscription background checker:", err)

        async def autopost_background_checker():
            while True:
                await asyncio.sleep(1800)  # فحص كل 30 دقيقة
                try:
                    from server import trigger_autopost_for_tenant
                    from datetime import datetime, timezone
                    active_autopost_tenants = (
                        supabase.table("tenants")
                        .select("id, name, autopost_enabled, autopost_interval_hours, autopost_last_at, facebook_page_id, facebook_page_token")
                        .eq("autopost_enabled", True)
                        .execute()
                        .data or []
                    )
                    now = datetime.now(timezone.utc)
                    for t in active_autopost_tenants:
                        if not t.get("facebook_page_id") or not t.get("facebook_page_token"):
                            continue
                        interval = int(t.get("autopost_interval_hours") or 24)
                        last_at_str = t.get("autopost_last_at")
                        should_post = False
                        if not last_at_str:
                            should_post = True
                        else:
                            try:
                                last_at = datetime.fromisoformat(str(last_at_str).replace("Z", "+00:00"))
                                diff_hours = (now - last_at).total_seconds() / 3600.0
                                if diff_hours >= interval:
                                    should_post = True
                            except Exception:
                                should_post = True

                        if should_post:
                            res = await trigger_autopost_for_tenant(t["id"])
                            if res.get("status") == "success":
                                print(f"📢 [Autopost] تم نشر بوست ترويجي تلقائي لمتجر «{t['name']}» بنجاح ({res.get('product_name')})", flush=True)
                except Exception as e:
                    print("Error in autopost background checker:", e)

        asyncio.create_task(background_checker())
        asyncio.create_task(autopost_background_checker())
        asyncio.create_task(check_subscriptions_and_notify(application))

    app = Application.builder().token(BOT_TOKEN).post_init(post_init).build()

    # تسجيل معالج الأخطاء لحماية الأمان
    app.add_error_handler(telegram_error_handler)

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(CommandHandler("stock", stock_cmd))
    app.add_handler(CommandHandler("inventory", stock_cmd))
    app.add_handler(CommandHandler("orders", orders_cmd))
    app.add_handler(CommandHandler("today", today_cmd))
    app.add_handler(CommandHandler("dashboard", dashboard_cmd))
    app.add_handler(CommandHandler("cancel", cancel_cmd))
    app.add_handler(CommandHandler("new_tenant", new_tenant_cmd))
    app.add_handler(CommandHandler("link_page", link_page_cmd))
    app.add_handler(CommandHandler("tenants", list_tenants_cmd))
    app.add_handler(CommandHandler("pause_bot", pause_bot_cmd))
    app.add_handler(CommandHandler("resume_bot", resume_bot_cmd))
    app.add_handler(CommandHandler("delete_tenant", delete_tenant_cmd))
    app.add_handler(CommandHandler("renew", renew_cmd))
    app.add_handler(CommandHandler("shipping", shipping_cmd))
    app.add_handler(CommandHandler("channels", channels_cmd))
    app.add_handler(CommandHandler("autopost", autopost_cmd))
    app.add_handler(CommandHandler("post_now", post_now_cmd))
    app.add_handler(CommandHandler("offers", offers_cmd))
    app.add_handler(CommandHandler("promos", offers_cmd))
    app.add_handler(CommandHandler("set_image", set_image_cmd))
    app.add_handler(CommandHandler("change_image", set_image_cmd))
    app.add_handler(CommandHandler("update_image", set_image_cmd))
    app.add_handler(CommandHandler("photo", set_image_cmd))
    app.add_handler(CallbackQueryHandler(on_button))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_handler(MessageHandler(filters.VOICE | filters.AUDIO, handle_voice))
    # الصور (عادية أو مبعوتة كملف) قبل معالج الملفات العام
    app.add_handler(MessageHandler(filters.PHOTO | filters.Document.IMAGE, handle_photo))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))

    print("بوت التاجر شغال... (Ctrl+C للإيقاف)", flush=True)
    app.run_polling()


if __name__ == "__main__":
    main()